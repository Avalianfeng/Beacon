#!/usr/bin/env python
"""提取上传文件的摘要 JSON，供前端展示和 analyst prompt 注入。

用法：
  python scripts/extract_file_meta.py <file_path> [purpose]
  purpose: problem | attachment（默认 attachment）

题面与文本类附件（PDF/DOCX/TXT/MD）共用 problem_ingest：
快速文本抽取、乱码质量、Markdown 落盘；PDF 乱码时标记 needsVision（视觉另调）。
Excel/CSV 仍走结构化摘要，不做视觉转写。
"""
from __future__ import annotations
import json
import sys
from pathlib import Path


def _meta_xlsx(path: Path) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for i, name in enumerate(wb.sheetnames):
        if i >= 5:
            sheets.append({"name": name, "rows": 0, "cols": 0, "columns": [], "preview": []})
            continue
        ws = wb[name]
        rows = list(ws.iter_rows(max_row=6, values_only=True))
        if not rows:
            sheets.append({"name": name, "rows": 0, "cols": 0, "columns": [], "preview": []})
            continue
        columns = [str(c) if c is not None else "" for c in rows[0]]
        preview = [[str(c) if c is not None else "" for c in row] for row in rows[1:6]]
        total_rows = sum(1 for _ in ws.iter_rows(values_only=True))
        sheets.append({
            "name": name, "rows": total_rows, "cols": len(columns),
            "columns": columns, "preview": preview,
        })
    wb.close()
    return {"sheets": sheets}


def _meta_csv(path: Path) -> dict:
    import csv
    with open(path, encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        rows = list(reader)
    if not rows:
        return {"sheets": [{"name": path.name, "rows": 0, "cols": 0, "columns": [], "preview": []}]}
    columns = [str(c) for c in rows[0]]
    preview = [[str(c) for c in row] for row in rows[1:6]]
    return {"sheets": [{"name": path.name, "rows": len(rows), "cols": len(columns),
                         "columns": columns, "preview": preview}]}


def _meta_ingest_text(
    path: Path,
    *,
    md_filename: str,
    text_limit: int,
) -> tuple[str, dict, dict]:
    """PDF/DOCX/TXT/MD：共用题面 ingest（不做视觉阻塞）。"""
    from math_agent.problem_ingest import parse_problem_file

    suffix = path.suffix.lower()
    result = parse_problem_file(
        path,
        enable_vision_fallback=False,
        write_md=True,
        md_filename=md_filename,
    )
    file_type = {
        ".pdf": "pdf",
        ".docx": "docx",
        ".txt": "txt",
        ".md": "txt",
    }.get(suffix, suffix.lstrip(".") or "txt")

    summary = result.to_summary_dict(text_limit=text_limit)
    if suffix == ".docx":
        from docx import Document
        doc = Document(str(path))
        summary["paragraphs"] = len([p for p in doc.paragraphs if p.text.strip()])
        summary["tables"] = len(doc.tables)
    if suffix in {".txt", ".md"}:
        summary["lines"] = result.text.count("\n") + 1

    extras = {
        "parsed_md_path": str(result.parsed_md_path).replace("\\", "/") if result.parsed_md_path else "",
        "parse_quality": result.quality.to_dict(),
        # 全文给题面文本框；summary.text_excerpt 仍截断供摘要/提示
        "text": result.text,
    }
    return file_type, summary, extras


def main():
    if len(sys.argv) < 2:
        print(json.dumps({"error": "usage: extract_file_meta.py <file_path> [purpose]"}))
        sys.exit(1)
    path = Path(sys.argv[1])
    purpose = (sys.argv[2] if len(sys.argv) > 2 else "attachment").strip().lower()
    if purpose not in {"problem", "attachment"}:
        purpose = "attachment"

    if not path.is_file():
        print(json.dumps({"error": f"file not found: {path}"}))
        sys.exit(1)

    suffix = path.suffix.lower()
    text_like = {".pdf", ".docx", ".txt", ".md"}

    if suffix in text_like:
        md_name = "problem_parsed.md" if purpose == "problem" else "attachment_parsed.md"
        limit = 5000 if suffix == ".pdf" else 3000
        try:
            file_type, summary, extras = _meta_ingest_text(
                path, md_filename=md_name, text_limit=limit,
            )
        except Exception as e:
            print(json.dumps({"error": f"extraction failed: {e}"}))
            sys.exit(1)
        print(json.dumps({
            "file_type": file_type,
            "filename": path.name,
            "summary": summary,
            "parsed_md_path": extras["parsed_md_path"],
            "parse_quality": extras["parse_quality"],
            "text": extras["text"],
        }, ensure_ascii=False))
        return

    type_map = {
        ".xlsx": ("xlsx", _meta_xlsx),
        ".xls": ("xlsx", _meta_xlsx),
        ".csv": ("csv", _meta_csv),
    }
    if suffix not in type_map:
        print(json.dumps({"error": f"unsupported file type: {suffix}"}))
        sys.exit(1)

    file_type, extractor = type_map[suffix]
    try:
        summary = extractor(path)
    except Exception as e:
        print(json.dumps({"error": f"extraction failed: {e}"}))
        sys.exit(1)

    print(json.dumps({
        "file_type": file_type,
        "filename": path.name,
        "summary": summary,
    }, ensure_ascii=False))


if __name__ == "__main__":
    for s in (sys.stdout, sys.stderr):
        if hasattr(s, "reconfigure"):
            s.reconfigure(encoding="utf-8", errors="replace")
    main()
