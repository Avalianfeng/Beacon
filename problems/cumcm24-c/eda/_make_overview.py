"""cumcm24-c EDA 出图（含图说 §3 重绘轮）。

附件均小，只读加载。不写建模结论。张数不限，特征优先。
"""
from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from matplotlib import font_manager
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "source"
OUT = Path(__file__).resolve().parent

TYPE_ORDER = ["平旱地", "梯田", "山坡地", "水浇地", "普通大棚", "智慧大棚"]
CROP_TYPE_ORDER = ["粮食（豆类）", "粮食", "蔬菜（豆类）", "蔬菜", "食用菌"]
BEAN_TYPES = {"粮食（豆类）", "蔬菜（豆类）"}


def _setup_cn_font() -> str:
    candidates = ("Microsoft YaHei", "SimHei", "Microsoft YaHei UI")
    available = {f.name for f in font_manager.fontManager.ttflist}
    chosen = next((n for n in candidates if n in available), "DejaVu Sans")
    plt.rcParams["font.sans-serif"] = [chosen, "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False
    return chosen


def _save(fig: plt.Figure, stem: str) -> Path:
    out = OUT / f"{stem}.png"
    fig.savefig(out, dpi=120, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return out


def _sheet_rows(path: Path, sheet: str) -> list[tuple]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    return rows


def _clean(s) -> str | None:
    if s is None:
        return None
    t = str(s).strip()
    return t or None


def load_plots() -> list[dict]:
    rows = _sheet_rows(SOURCE / "附件1.xlsx", "乡村的现有耕地")
    plots = []
    for r in rows[1:]:
        name = _clean(r[0])
        if not name or name.startswith("(") or name.startswith("注"):
            continue
        typ = _clean(r[1])
        if typ is None:
            continue
        plots.append({"name": name, "type": typ, "area": float(r[2]) if r[2] is not None else None})
    return plots


def load_crops() -> list[dict]:
    rows = _sheet_rows(SOURCE / "附件1.xlsx", "乡村种植的农作物")
    crops = []
    for r in rows[1:]:
        if r[0] is None or (isinstance(r[0], str) and not str(r[0]).isdigit()):
            continue
        crops.append(
            {
                "id": int(r[0]),
                "name": _clean(r[1]),
                "type": _clean(r[2]),
            }
        )
    return crops


def load_plant2023() -> list[dict]:
    rows = _sheet_rows(SOURCE / "附件2.xlsx", "2023年的农作物种植情况")
    plant, last_plot = [], None
    for r in rows[1:]:
        if r[1] is None:
            continue
        plot = _clean(r[0]) or last_plot
        if plot:
            last_plot = plot
        plant.append(
            {
                "plot": plot,
                "crop_id": int(r[1]),
                "crop": _clean(r[2]),
                "type": _clean(r[3]),
                "area": float(r[4]) if r[4] is not None else None,
                "season": _clean(r[5]),
            }
        )
    return plant


def _parse_price(p) -> dict | None:
    if p is None:
        return None
    if isinstance(p, (int, float)):
        return {"lo": float(p), "hi": float(p), "raw": str(p)}
    s = str(p).strip()
    m = re.match(r"^([0-9.]+)\s*-\s*([0-9.]+)$", s)
    if m:
        return {"lo": float(m.group(1)), "hi": float(m.group(2)), "raw": s}
    try:
        v = float(s)
        return {"lo": v, "hi": v, "raw": s}
    except ValueError:
        return {"lo": None, "hi": None, "raw": s}


def load_stats() -> list[dict]:
    rows = _sheet_rows(SOURCE / "附件2.xlsx", "2023年统计的相关数据")
    out = []
    for r in rows[1:]:
        if r[1] is None:
            continue
        try:
            cid = int(r[1])
        except (TypeError, ValueError):
            continue
        out.append(
            {
                "crop_id": cid,
                "crop": _clean(r[2]),
                "land_type": _clean(r[3]),
                "season": _clean(r[4]),
                "yield_jin_per_mu": float(r[5]) if isinstance(r[5], (int, float)) else None,
                "cost_yuan_per_mu": float(r[6]) if isinstance(r[6], (int, float)) else None,
                "price": _parse_price(r[7]),
            }
        )
    return out


def profile_result_template(path: Path) -> dict:
    wb = load_workbook(path, read_only=True, data_only=True)
    sheets = list(wb.sheetnames)
    rows = list(wb[sheets[0]].iter_rows(values_only=True))
    header = rows[0]
    crops = [c for c in header[2:] if c]
    plots_s1, plots_s2 = [], []
    season = "第一季"
    for r in rows[1:]:
        marker = _clean(r[0])
        if marker and "二" in marker.replace("\n", ""):
            season = "第二季"
        if marker and marker.startswith("注"):
            break
        plot = _clean(r[1])
        if not plot or not re.match(r"^[A-Z]", plot):
            continue
        (plots_s2 if "二" in season else plots_s1).append(plot)
    nonempty = 0
    for sh in sheets:
        for row in wb[sh].iter_rows(min_row=2, min_col=3, values_only=True):
            for v in row:
                if v is not None and str(v).strip() not in {"", "None"}:
                    nonempty += 1
    wb.close()
    return {
        "file": path.name,
        "sheets": sheets,
        "n_crop_cols": len(crops),
        "crop_cols": [str(c).strip() for c in crops],
        "plots_season1": plots_s1,
        "plots_season2": plots_s2,
        "n_plots_s1": len(plots_s1),
        "n_plots_s2": len(plots_s2),
        "nonempty_value_cells": nonempty,
    }


def make_attachment1(plots: list[dict], crops: list[dict]) -> Path:
    by_t = defaultdict(list)
    for p in plots:
        by_t[p["type"]].append(p)
    areas = [sum(x["area"] or 0 for x in by_t[t]) for t in TYPE_ORDER]
    counts = [len(by_t[t]) for t in TYPE_ORDER]
    crop_counts = Counter(c["type"] for c in crops)

    fig, axes = plt.subplots(1, 2, figsize=(12, 5))
    x = np.arange(len(TYPE_ORDER))
    bars = axes[0].bar(x, areas, color="#4C78A8")
    axes[0].set_xticks(x)
    axes[0].set_xticklabels(TYPE_ORDER, rotation=25, ha="right")
    axes[0].set_ylabel("面积合计（亩）")
    axes[0].set_title(f"附件1·耕地：{len(plots)} 块 / 合计 {sum(areas):.1f} 亩")
    for b, n in zip(bars, counts):
        axes[0].text(b.get_x() + b.get_width() / 2, b.get_height(), f"n={n}", ha="center", va="bottom", fontsize=8)

    cy = [crop_counts.get(t, 0) for t in CROP_TYPE_ORDER]
    axes[1].barh(CROP_TYPE_ORDER, cy, color="#F58518")
    axes[1].set_xlabel("作物种数")
    axes[1].set_title(f"附件1·作物：共 {len(crops)} 种")
    for i, v in enumerate(cy):
        axes[1].text(v + 0.1, i, str(v), va="center", fontsize=9)
    fig.tight_layout()
    return _save(fig, "附件1")


def make_plot_area_spread(plots: list[dict]) -> Path:
    """§3.1 同地类内地块面积离散（点图）。"""
    fig, ax = plt.subplots(figsize=(11, 5))
    rng = np.random.default_rng(0)
    for i, t in enumerate(TYPE_ORDER):
        vals = [p["area"] for p in plots if p["type"] == t and p["area"] is not None]
        if not vals:
            continue
        jitter = rng.uniform(-0.18, 0.18, size=len(vals))
        ax.scatter(np.full(len(vals), i) + jitter, vals, s=36, alpha=0.85, zorder=3)
        ax.hlines(np.median(vals), i - 0.25, i + 0.25, colors="black", lw=1.5, zorder=4)
        for p in plots:
            if p["type"] == t and p["area"] is not None and (p["area"] == max(vals) or p["area"] == min(vals)):
                ax.annotate(p["name"], (i, p["area"]), textcoords="offset points", xytext=(4, 2), fontsize=7)
    ax.set_xticks(range(len(TYPE_ORDER)))
    ax.set_xticklabels(TYPE_ORDER, rotation=20, ha="right")
    ax.set_ylabel("单块面积（亩）")
    ax.set_title("地块面积离散（同地类内；横线=中位数）")
    ax.grid(axis="y", alpha=0.3)
    fig.tight_layout()
    return _save(fig, "补-地块面积离散")


def make_attachment2(plant: list[dict], stats: list[dict]) -> Path:
    type_area = Counter()
    season_area = Counter()
    crop_area = Counter()
    for p in plant:
        type_area[p["type"] or "?"] += p["area"] or 0
        season_area[p["season"] or "?"] += p["area"] or 0
        crop_area[p["crop"] or str(p["crop_id"])] += p["area"] or 0
    land_types = sorted({s["land_type"] for s in stats if s["land_type"]})

    fig, axes = plt.subplots(2, 2, figsize=(12, 8))
    labels = list(type_area.keys())
    vals = [type_area[k] for k in labels]
    axes[0, 0].bar(labels, vals, color="#54A24B")
    axes[0, 0].set_title(f"2023 种植面积按作物类型（行合计 {sum(vals):.1f} 亩）")
    axes[0, 0].tick_params(axis="x", rotation=20)
    axes[0, 0].set_ylabel("亩")

    sl = list(season_area.keys())
    axes[0, 1].bar(sl, [season_area[k] for k in sl], color="#E45756")
    axes[0, 1].set_title("2023 种植面积按季次")
    axes[0, 1].set_ylabel("亩")

    top = crop_area.most_common(15)
    axes[1, 0].barh([t[0] for t in reversed(top)], [t[1] for t in reversed(top)], color="#B279A2")
    axes[1, 0].set_title("2023 种植面积 Top15 作物")
    axes[1, 0].set_xlabel("亩")

    colors = plt.cm.tab10(np.linspace(0, 1, max(len(land_types), 1)))
    for i, lt in enumerate(land_types):
        xs = [s["cost_yuan_per_mu"] for s in stats if s["land_type"] == lt and s["cost_yuan_per_mu"] is not None]
        ys = [s["yield_jin_per_mu"] for s in stats if s["land_type"] == lt and s["yield_jin_per_mu"] is not None]
        axes[1, 1].scatter(xs, ys, s=28, alpha=0.75, label=lt, color=colors[i])
    axes[1, 1].set_xlabel("种植成本（元/亩）")
    axes[1, 1].set_ylabel("亩产量（斤/亩）")
    axes[1, 1].set_title(f"统计表：成本×亩产（{len(stats)} 行）")
    axes[1, 1].set_xscale("log")
    axes[1, 1].set_yscale("log")
    axes[1, 1].legend(fontsize=7, loc="lower right")
    fig.tight_layout()
    return _save(fig, "附件2")


def make_utilization(plots: list[dict], plant: list[dict]) -> Path:
    """§3.2 行面积合计 / 地块面积（>1 表示两季或合种复计）。"""
    plot_area = {p["name"]: p["area"] for p in plots}
    plot_type = {p["name"]: p["type"] for p in plots}
    row_sum = Counter()
    for p in plant:
        if p["plot"]:
            row_sum[p["plot"]] += p["area"] or 0

    names = sorted(plot_area.keys(), key=lambda n: (TYPE_ORDER.index(plot_type[n]) if plot_type[n] in TYPE_ORDER else 99, n))
    util = []
    for n in names:
        a = plot_area[n] or 0
        util.append((row_sum[n] / a) if a else 0.0)

    colors = []
    for n in names:
        t = plot_type[n]
        colors.append(plt.cm.Set2(TYPE_ORDER.index(t) / max(len(TYPE_ORDER) - 1, 1)) if t in TYPE_ORDER else "#999")

    fig, ax = plt.subplots(figsize=(14, 5))
    x = np.arange(len(names))
    ax.bar(x, util, color=colors, width=0.85)
    ax.axhline(1.0, color="black", ls="--", lw=1, label="利用率=1（单季满种）")
    ax.set_xticks(x)
    ax.set_xticklabels(names, rotation=90, fontsize=7)
    ax.set_ylabel("种植行面积合计 / 地块面积")
    ax.set_title("2023 各地块「行面积/地块面积」（>1 = 两季或合种复计）")
    ax.legend(loc="upper right", fontsize=8)
    ax.set_ylim(0, max(util) * 1.15 if util else 1)
    fig.tight_layout()
    return _save(fig, "补-地块利用率")


def make_price_ranges(stats: list[dict], crops: list[dict]) -> Path:
    """§3.3 售价区间：按作物取各行 lo/hi 包络。"""
    id_to_type = {c["id"]: c["type"] for c in crops}
    by_crop: dict[int, dict] = {}
    for s in stats:
        pr = s["price"]
        if not pr or pr["lo"] is None or pr["hi"] is None:
            continue
        cid = s["crop_id"]
        slot = by_crop.setdefault(
            cid,
            {"name": s["crop"], "type": id_to_type.get(cid, "?"), "lo": pr["lo"], "hi": pr["hi"]},
        )
        slot["lo"] = min(slot["lo"], pr["lo"])
        slot["hi"] = max(slot["hi"], pr["hi"])

    # sort by type then mid
    items = sorted(
        by_crop.values(),
        key=lambda d: (
            CROP_TYPE_ORDER.index(d["type"]) if d["type"] in CROP_TYPE_ORDER else 99,
            (d["lo"] + d["hi"]) / 2,
        ),
    )
    fig, ax = plt.subplots(figsize=(10, 10))
    y = np.arange(len(items))
    type_color = {t: plt.cm.tab10(i) for i, t in enumerate(CROP_TYPE_ORDER)}
    for i, d in enumerate(items):
        c = type_color.get(d["type"], "#888")
        ax.hlines(i, d["lo"], d["hi"], colors=c, lw=3)
        ax.plot([d["lo"], d["hi"]], [i, i], "o", color=c, markersize=4)
    ax.set_yticks(y)
    ax.set_yticklabels([d["name"] for d in items], fontsize=8)
    ax.set_xlabel("销售单价（元/斤）")
    ax.set_title("统计表售价区间包络（同作物跨地类取 min lo–max hi；无销量列）")
    ax.grid(axis="x", alpha=0.3)
    # legend
    for t in CROP_TYPE_ORDER:
        ax.plot([], [], color=type_color[t], lw=3, label=t)
    ax.legend(loc="lower right", fontsize=8)
    fig.tight_layout()
    return _save(fig, "补-售价区间")


def make_attachment3(tmpl: dict) -> Path:
    fig, axes = plt.subplots(1, 2, figsize=(12, 5))
    years = tmpl["sheets"]
    axes[0].bar(years, [1] * len(years), color="#72B7B2")
    axes[0].set_ylim(0, 1.5)
    axes[0].set_title(f"result*.xlsx 年 sheet（共 {len(years)}）")
    axes[0].set_yticks([])
    cats = ["第一季地块行", "第二季地块行", "作物列"]
    vals = [tmpl["n_plots_s1"], tmpl["n_plots_s2"], tmpl["n_crop_cols"]]
    axes[1].bar(cats, vals, color=["#4C78A8", "#F58518", "#54A24B"])
    axes[1].set_title(f"{tmpl['file']} 结构（填值非空={tmpl['nonempty_value_cells']}）")
    for i, v in enumerate(vals):
        axes[1].text(i, v, str(v), ha="center", va="bottom")
    fig.suptitle("附件3：提交模板（空表）结构示意", y=1.02)
    fig.tight_layout()
    return _save(fig, "附件3")


def make_season_coverage(tmpl: dict, plant: list[dict]) -> Path:
    """§3.4 模板第二季行 vs 2023 是否实际种第二季。"""
    s2_plots = tmpl["plots_season2"]
    planted_s2 = {p["plot"] for p in plant if p["season"] == "第二季" and p["plot"]}
    yes = [1 if n in planted_s2 else 0 for n in s2_plots]
    colors = ["#54A24B" if y else "#E45756" for y in yes]

    fig, ax = plt.subplots(figsize=(12, 4.5))
    x = np.arange(len(s2_plots))
    ax.bar(x, [1] * len(s2_plots), color=colors, width=0.85)
    ax.set_xticks(x)
    ax.set_xticklabels(s2_plots, rotation=90, fontsize=8)
    ax.set_yticks([])
    ax.set_title("模板第二季行集：绿=2023 实际有第二季；红=模板有行但 2023 无第二季（如 D7/D8）")
    # annotate red ones
    for i, (n, y) in enumerate(zip(s2_plots, yes)):
        if not y:
            ax.text(i, 1.02, n, ha="center", va="bottom", fontsize=8, color="#E45756")
    n_miss = sum(1 for y in yes if not y)
    ax.text(
        0.99,
        0.95,
        f"模板第二季 {len(s2_plots)} 行；2023 缺第二季 {n_miss} 块",
        transform=ax.transAxes,
        ha="right",
        va="top",
        fontsize=9,
        bbox=dict(boxstyle="round", facecolor="white", alpha=0.8),
    )
    fig.tight_layout()
    return _save(fig, "补-模板季次覆盖")


def make_bean_baseline(plots: list[dict], plant: list[dict]) -> Path:
    """§3.5 2023 豆类落点（粮豆+蔬豆）；不断言三年策略。"""
    bean_by_plot: dict[str, list[str]] = defaultdict(list)
    for p in plant:
        if p["type"] in BEAN_TYPES and p["plot"]:
            bean_by_plot[p["plot"]].append(f"{p['crop']}({p['season']})")

    # grid by type
    fig, axes = plt.subplots(2, 3, figsize=(12, 7), sharey=False)
    axes_flat = axes.ravel()
    for i, t in enumerate(TYPE_ORDER):
        ax = axes_flat[i]
        names = [p["name"] for p in plots if p["type"] == t]
        has = [1 if n in bean_by_plot else 0 for n in names]
        colors = ["#4C78A8" if h else "#DDD" for h in has]
        ax.barh(names[::-1], [1] * len(names), color=colors[::-1])
        ax.set_xlim(0, 1.2)
        ax.set_xticks([])
        ax.set_title(f"{t}（蓝=2023种过豆类）", fontsize=10)
        for j, n in enumerate(names[::-1]):
            if n in bean_by_plot:
                ax.text(1.02, j, ",".join(bean_by_plot[n][:2]), va="center", fontsize=6)
    n_bean_plots = len(bean_by_plot)
    fig.suptitle(f"2023 豆类基线落点：{n_bean_plots}/{len(plots)} 块曾种粮豆或蔬豆（非策略建议）", y=1.01)
    fig.tight_layout()
    return _save(fig, "补-豆类落点")


def main() -> None:
    font = _setup_cn_font()
    plots = load_plots()
    crops = load_crops()
    plant = load_plant2023()
    stats = load_stats()
    tmpl = profile_result_template(SOURCE / "result1_1.xlsx")
    tmpl2 = profile_result_template(SOURCE / "result1_2.xlsx")
    tmpl3 = profile_result_template(SOURCE / "result2.xlsx")

    by_type = defaultdict(list)
    for p in plots:
        by_type[p["type"]].append(p)
    plot_summary = {
        t: {"n": len(lst), "area": round(sum(x["area"] or 0 for x in lst), 2), "names": [x["name"] for x in lst]}
        for t, lst in by_type.items()
    }
    plant_by_plot = defaultdict(list)
    for p in plant:
        plant_by_plot[p["plot"]].append(p)

    # utilization stats
    util_map = {}
    for p in plots:
        a = p["area"] or 0
        util_map[p["name"]] = round(sum(x["area"] or 0 for x in plant_by_plot[p["name"]]) / a, 3) if a else None
    s2_set = set(tmpl["plots_season2"])
    planted_s2 = {p["plot"] for p in plant if p["season"] == "第二季"}
    s2_missing = sorted(s2_set - planted_s2)
    bean_plots = sorted({p["plot"] for p in plant if p["type"] in BEAN_TYPES and p["plot"]})

    figures = [
        make_attachment1(plots, crops),
        make_plot_area_spread(plots),
        make_attachment2(plant, stats),
        make_utilization(plots, plant),
        make_price_ranges(stats, crops),
        make_attachment3(tmpl),
        make_season_coverage(tmpl, plant),
        make_bean_baseline(plots, plant),
    ]

    stats_out = {
        "font": font,
        "round": "图说§3重绘",
        "plots": {
            "n": len(plots),
            "open_area_mu": sum(plot_summary[t]["area"] for t in ("平旱地", "梯田", "山坡地", "水浇地") if t in plot_summary),
            "total_area_mu": round(sum(p["area"] or 0 for p in plots), 2),
            "by_type": plot_summary,
        },
        "crops": {"n": len(crops), "by_type": dict(Counter(c["type"] for c in crops))},
        "plant_2023": {
            "n_rows": len(plant),
            "total_area_rows_mu": round(sum(p["area"] or 0 for p in plant), 2),
            "season_counts": dict(Counter(p["season"] for p in plant)),
            "type_area_mu": {
                t: round(sum(p["area"] or 0 for p in plant if p["type"] == t), 2)
                for t in sorted({p["type"] for p in plant})
            },
        },
        "stats_2023": {
            "n_rows": len(stats),
            "no_expected_sales_column": True,
            "n_price_range_strings": len({s["price"]["raw"] for s in stats if s["price"] and s["price"]["raw"]}),
        },
        "utilization": {
            "min": min(v for v in util_map.values() if v is not None),
            "max": max(v for v in util_map.values() if v is not None),
            "gt1_plots": sorted(k for k, v in util_map.items() if v is not None and v > 1.001),
            "note": ">1 表示行面积合计超过地块面积（两季/合种复计）",
        },
        "season_coverage": {
            "template_s2_n": len(s2_set),
            "planted_s2_n": len(planted_s2 & s2_set),
            "template_s2_missing_in_2023": s2_missing,
        },
        "bean_baseline": {"n_plots_with_bean": len(bean_plots), "plots": bean_plots},
        "result_templates": {
            "result1_1": {k: tmpl[k] for k in ("file", "sheets", "n_plots_s1", "n_plots_s2", "n_crop_cols", "nonempty_value_cells")},
            "same_structure": tmpl2["n_plots_s1"] == tmpl["n_plots_s1"] and tmpl3["n_plots_s1"] == tmpl["n_plots_s1"],
        },
        "figures": [p.name for p in figures],
    }
    (OUT / "_profile_stats.json").write_text(json.dumps(stats_out, ensure_ascii=False, indent=2), encoding="utf-8")
    print("font:", font)
    for p in figures:
        print("wrote", p.name, p.stat().st_size)
    print("util range", stats_out["utilization"]["min"], stats_out["utilization"]["max"])
    print("s2 missing", s2_missing)
    print("bean plots", len(bean_plots))


if __name__ == "__main__":
    main()
