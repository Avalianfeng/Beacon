"""兑现第二家 8 条索取：附件2 只 openpyxl read_only 流式扫描一次。"""
from __future__ import annotations

import calendar
import math
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib import font_manager
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "source"
OUT = Path(__file__).resolve().parent

DATA_START = date(2020, 7, 1)
DATA_END = date(2023, 6, 30)
Q3_START = date(2023, 6, 24)
Q3_END = date(2023, 6, 30)
WIN_END = date(2023, 6, 30)

CAT_ORDER = ["花叶类", "花菜类", "水生根茎类", "茄类", "辣椒类", "食用菌"]


def _setup_cn_font() -> str:
    candidates = ("Microsoft YaHei", "SimHei", "Microsoft YaHei UI")
    available = {f.name for f in font_manager.fontManager.ttflist}
    chosen = next((n for n in candidates if n in available), "DejaVu Sans")
    plt.rcParams["font.sans-serif"] = [chosen, "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False
    return chosen


def _norm_code(x) -> str | None:
    if x is None:
        return None
    if isinstance(x, float):
        if np.isnan(x):
            return None
        if x.is_integer():
            return str(int(x))
        return str(x).strip()
    if isinstance(x, int):
        return str(x)
    s = str(x).strip()
    return s or None


def _as_date(x) -> date | None:
    if x is None:
        return None
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    try:
        return pd.Timestamp(x).date()
    except Exception:
        return None


def _pct(arr: np.ndarray, q: float) -> float:
    if arr.size == 0:
        return float("nan")
    return float(np.percentile(arr, q))


def _std_sample(arr: np.ndarray) -> float:
    if arr.size < 2:
        return float("nan")
    return float(np.std(arr, ddof=1))


def _mean(arr: np.ndarray) -> float:
    if arr.size == 0:
        return float("nan")
    return float(np.mean(arr))


def _round(x: float, n: int = 6) -> float | None:
    if x is None or (isinstance(x, float) and (math.isnan(x) or math.isinf(x))):
        return None
    return round(float(x), n)


def _pearson(n: int, sx: float, sy: float, sxy: float, sx2: float, sy2: float) -> float:
    if n < 10:
        return float("nan")
    num = n * sxy - sx * sy
    den_x = n * sx2 - sx * sx
    den_y = n * sy2 - sy * sy
    if den_x <= 0 or den_y <= 0:
        return float("nan")
    den = math.sqrt(den_x * den_y)
    if den == 0:
        return float("nan")
    return num / den


def daterange(start: date, end: date):
    d = start
    while d <= end:
        yield d
        d += timedelta(days=1)


def calendar_days_in_month(year: int, month: int, clip_start: date, clip_end: date) -> list[date]:
    last = calendar.monthrange(year, month)[1]
    a = date(year, month, 1)
    b = date(year, month, last)
    a = max(a, clip_start)
    b = min(b, clip_end)
    return list(daterange(a, b))


def load_att1() -> tuple[pd.DataFrame, dict[str, str], dict[str, str], dict[str, int]]:
    df = pd.read_excel(SOURCE / "附件1.xlsx", sheet_name="Sheet1", dtype=object)
    df.columns = [str(c).strip() for c in df.columns]
    df["单品编码"] = df["单品编码"].map(_norm_code)
    sku_to_cat = dict(zip(df["单品编码"], df["分类名称"].astype(str)))
    sku_to_name = dict(zip(df["单品编码"], df["单品名称"].astype(str)))
    cat_n_sku = df["分类名称"].value_counts().to_dict()
    return df, sku_to_cat, sku_to_name, cat_n_sku


def load_att3() -> dict[tuple[date, str], float]:
    df = pd.read_excel(SOURCE / "附件3.xlsx", sheet_name="Sheet1", dtype=object)
    df.columns = [str(c).strip() for c in df.columns]
    df["单品编码"] = df["单品编码"].map(_norm_code)
    out: dict[tuple[date, str], float] = {}
    for raw_d, sku, raw_p in zip(df["日期"], df["单品编码"], df["批发价格(元/千克)"]):
        d = _as_date(raw_d)
        if d is None or not sku:
            continue
        try:
            p = float(raw_p)
        except (TypeError, ValueError):
            continue
        if math.isnan(p):
            continue
        out[(d, sku)] = p
    return out


def load_att4_sku() -> pd.DataFrame:
    xl = pd.ExcelFile(SOURCE / "附件4.xlsx")
    sheets = list(xl.sheet_names)
    cat_sheet = next(s for s in sheets if "损耗率" in s and "分类" in s)
    sku_sheet = "Sheet1" if "Sheet1" in sheets else [s for s in sheets if s != cat_sheet][0]
    df = pd.read_excel(xl, sheet_name=sku_sheet, dtype=object)
    df.columns = [str(c).strip() for c in df.columns]
    df["单品编码"] = df["单品编码"].map(_norm_code)
    df["损耗率(%)"] = pd.to_numeric(df["损耗率(%)"], errors="coerce")
    return df


class PearsonAcc:
    __slots__ = ("n", "sx", "sy", "sxy", "sx2", "sy2")

    def __init__(self) -> None:
        self.n = 0
        self.sx = 0.0
        self.sy = 0.0
        self.sxy = 0.0
        self.sx2 = 0.0
        self.sy2 = 0.0

    def add(self, x: float, y: float) -> None:
        self.n += 1
        self.sx += x
        self.sy += y
        self.sxy += x * y
        self.sx2 += x * x
        self.sy2 += y * y

    def r(self) -> float:
        return _pearson(self.n, self.sx, self.sy, self.sxy, self.sx2, self.sy2)


def stream_att2(
    sku_to_cat: dict[str, str],
) -> dict:
    p = SOURCE / "附件2.xlsx"
    wb = load_workbook(p, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    header = [h if h is None else str(h).strip() for h in next(it)]
    idx = {name: i for i, name in enumerate(header)}
    required = ["销售日期", "单品编码", "销量(千克)", "销售单价(元/千克)", "销售类型", "是否打折销售"]
    missing_cols = [c for c in required if c not in idx]
    if missing_cols:
        wb.close()
        raise KeyError(f"附件2 缺列: {missing_cols}; header={header}")

    i_date = idx["销售日期"]
    i_sku = idx["单品编码"]
    i_qty = idx["销量(千克)"]
    i_price = idx["销售单价(元/千克)"]
    i_type = idx["销售类型"]
    i_disc = idx["是否打折销售"]

    daily_cat_qty: dict[date, dict[str, float]] = defaultdict(lambda: defaultdict(float))
    cat_sale_days: dict[str, set[date]] = defaultdict(set)
    cat_skus_flow: dict[str, set[str]] = defaultdict(set)
    pearson: dict[tuple[str, str, str], PearsonAcc] = defaultdict(PearsonAcc)
    sold_pairs: set[tuple[date, str]] = set()
    q3_sku_qty: dict[str, float] = defaultdict(float)
    q3_sku_days: dict[str, set[date]] = defaultdict(set)
    ret_cat_rows: dict[str, int] = defaultdict(int)
    ret_cat_qty: dict[str, float] = defaultdict(float)
    sale_cat_qty: dict[str, float] = defaultdict(float)
    ret_sku_rows: dict[str, int] = defaultdict(int)
    ret_sku_qty: dict[str, float] = defaultdict(float)

    n = 0
    for raw in it:
        if raw is None or all(v is None for v in raw):
            continue
        n += 1
        d = _as_date(raw[i_date] if i_date < len(raw) else None)
        sku = _norm_code(raw[i_sku] if i_sku < len(raw) else None)
        stype = str(raw[i_type]).strip() if raw[i_type] is not None else ""
        disc = str(raw[i_disc]).strip() if raw[i_disc] is not None else ""
        try:
            qv = float(raw[i_qty]) if raw[i_qty] is not None else None
        except (TypeError, ValueError):
            qv = None
        try:
            pv = float(raw[i_price]) if raw[i_price] is not None else None
        except (TypeError, ValueError):
            pv = None

        cat = sku_to_cat.get(sku, "未知") if sku else "未知"

        if d is not None and qv is not None:
            daily_cat_qty[d][cat] += qv
            cat_sale_days[cat].add(d)
        if sku:
            cat_skus_flow[cat].add(sku)

        if qv is not None and pv is not None and d is not None:
            ym = f"{d.year:04d}-{d.month:02d}"
            pearson[(ym, cat, "全部")].add(qv, pv)
            if disc:
                pearson[(ym, cat, disc)].add(qv, pv)

        if d is not None and sku and stype == "销售":
            sold_pairs.add((d, sku))

        if d is not None and sku and Q3_START <= d <= Q3_END:
            q3_sku_days[sku].add(d)
            if qv is not None:
                q3_sku_qty[sku] += qv

        if stype == "退货":
            ret_cat_rows[cat] += 1
            if qv is not None:
                ret_cat_qty[cat] += qv
            if sku:
                ret_sku_rows[sku] += 1
                if qv is not None:
                    ret_sku_qty[sku] += qv
        elif stype == "销售":
            if qv is not None:
                sale_cat_qty[cat] += qv

        if n % 200_000 == 0:
            print(f"  附件2 已扫描 {n:,} 行…", flush=True)

    wb.close()
    print(f"  附件2 扫描完成 {n:,} 行", flush=True)
    return {
        "n": n,
        "daily_cat_qty": daily_cat_qty,
        "cat_sale_days": cat_sale_days,
        "cat_skus_flow": cat_skus_flow,
        "pearson": pearson,
        "sold_pairs": sold_pairs,
        "q3_sku_qty": q3_sku_qty,
        "q3_sku_days": q3_sku_days,
        "ret_cat_rows": ret_cat_rows,
        "ret_cat_qty": ret_cat_qty,
        "sale_cat_qty": sale_cat_qty,
        "ret_sku_rows": ret_sku_rows,
        "ret_sku_qty": ret_sku_qty,
    }


def daily_series(daily_cat_qty, cat: str, start: date, end: date, fill0: bool) -> np.ndarray:
    if fill0:
        vals = [float(daily_cat_qty.get(d, {}).get(cat, 0.0)) for d in daterange(start, end)]
        return np.asarray(vals, dtype=float)
    vals = []
    d = start
    while d <= end:
        if d in daily_cat_qty:
            vals.append(float(daily_cat_qty[d].get(cat, 0.0)))
        d += timedelta(days=1)
    return np.asarray(vals, dtype=float)


def write_01(daily_cat_qty, cat_sale_days, cats: list[str]) -> Path:
    rows = []
    months: set[tuple[int, int]] = set()
    for d in daily_cat_qty:
        months.add((d.year, d.month))
    for year, month in sorted(months):
        days = calendar_days_in_month(year, month, DATA_START, DATA_END)
        ym = f"{year:04d}-{month:02d}"
        n_cal = len(days)
        for cat in cats:
            qty_by_day = np.asarray(
                [float(daily_cat_qty.get(d, {}).get(cat, 0.0)) for d in days], dtype=float
            )
            total = float(qty_by_day.sum())
            sale_days = {d for d in cat_sale_days.get(cat, set()) if d.year == year and d.month == month}
            n_sale = len(sale_days)
            rows.append(
                {
                    "年月": ym,
                    "品类": cat,
                    "月净销量合计_kg": _round(total, 6),
                    "日历天数": n_cal,
                    "有销售天数": n_sale,
                    "日均_按日历天数_kg": _round(total / n_cal, 6) if n_cal else None,
                    "日均_按有销售天数_kg": _round(total / n_sale, 6) if n_sale else None,
                    "日销量均值_日历补0_kg": _round(_mean(qty_by_day), 6),
                    "日销量中位数_日历补0_kg": _round(float(np.median(qty_by_day)), 6) if n_cal else None,
                    "日销量P25_日历补0_kg": _round(_pct(qty_by_day, 25), 6),
                    "日销量P75_日历补0_kg": _round(_pct(qty_by_day, 75), 6),
                    "日销量标准差_日历补0_kg": _round(_std_sample(qty_by_day), 6),
                }
            )
    out = OUT / "01-月品类净销量分布.csv"
    pd.DataFrame(rows).to_csv(out, index=False, encoding="utf-8-sig")
    return out


def write_02(daily_cat_qty, cat_skus_flow, cat_n_sku, cats: list[str]) -> tuple[Path, dict[str, np.ndarray]]:
    n_cal = (DATA_END - DATA_START).days + 1
    rows = []
    series = {}
    for cat in cats:
        arr = np.asarray(
            [float(daily_cat_qty.get(d, {}).get(cat, 0.0)) for d in daterange(DATA_START, DATA_END)],
            dtype=float,
        )
        series[cat] = arr
        n_zero = int(np.sum(np.abs(arr) < 1e-12))
        rows.append(
            {
                "品类": cat,
                "日历天数": n_cal,
                "零值天数": n_zero,
                "零值比例": _round(n_zero / n_cal, 6) if n_cal else None,
                "P50_kg": _round(_pct(arr, 50), 6),
                "P75_kg": _round(_pct(arr, 75), 6),
                "P90_kg": _round(_pct(arr, 90), 6),
                "P95_kg": _round(_pct(arr, 95), 6),
                "P99_kg": _round(_pct(arr, 99), 6),
                "最大值_kg": _round(float(arr.max()) if arr.size else float("nan"), 6),
                "目录SKU数": int(cat_n_sku.get(cat, 0)),
                "流水出现SKU数": len(cat_skus_flow.get(cat, set())),
            }
        )
    out = OUT / "02-品类日销量分位.csv"
    pd.DataFrame(rows).to_csv(out, index=False, encoding="utf-8-sig")
    return out, series


def write_03(pearson, cats: list[str]) -> Path:
    keys = set(pearson)
    yms = sorted({k[0] for k in keys})
    rows = []
    for ym in yms:
        for cat in cats:
            a_no = pearson.get((ym, cat, "否"))
            a_yes = pearson.get((ym, cat, "是"))
            a_all = pearson.get((ym, cat, "全部"))
            n_no = a_no.n if a_no else 0
            n_yes = a_yes.n if a_yes else 0
            n_all = a_all.n if a_all else 0
            r_no = a_no.r() if a_no else float("nan")
            r_yes = a_yes.r() if a_yes else float("nan")
            r_all = a_all.r() if a_all else float("nan")
            rows.append(
                {
                    "年月": ym,
                    "品类": cat,
                    "n_否": n_no,
                    "pearson_否": None if math.isnan(r_no) else _round(r_no, 6),
                    "n_是": n_yes,
                    "pearson_是": None if math.isnan(r_yes) else _round(r_yes, 6),
                    "n_不分打折": n_all,
                    "pearson_不分打折": None if math.isnan(r_all) else _round(r_all, 6),
                }
            )
    out = OUT / "03-品类月-量价相关.csv"
    pd.DataFrame(rows).to_csv(out, index=False, encoding="utf-8-sig")
    return out


def write_04(
    sold_pairs: set[tuple[date, str]],
    wh: dict[tuple[date, str], float],
    sku_to_cat: dict[str, str],
    sku_to_name: dict[str, str],
    cats: list[str],
) -> tuple[Path, Path, dict[str, np.ndarray]]:
    cat_prices: dict[str, list[float]] = defaultdict(list)
    cat_ext: dict[str, list[tuple[float, date, str]]] = defaultdict(list)
    n_hit = 0
    n_miss = 0
    for d, sku in sold_pairs:
        price = wh.get((d, sku))
        if price is None:
            n_miss += 1
            continue
        n_hit += 1
        cat = sku_to_cat.get(sku, "未知")
        cat_prices[cat].append(price)
        cat_ext[cat].append((price, d, sku))
    rows = []
    series = {}
    extra_rows = []
    for cat in cats + (["未知"] if cat_prices.get("未知") else []):
        arr = np.asarray(cat_prices.get(cat, []), dtype=float)
        series[cat] = arr
        if arr.size == 0:
            rows.append({"品类": cat, "n_已售且有批发价": 0})
            continue
        rows.append(
            {
                "品类": cat,
                "n_已售且有批发价": int(arr.size),
                "P1": _round(_pct(arr, 1), 6),
                "P5": _round(_pct(arr, 5), 6),
                "P25": _round(_pct(arr, 25), 6),
                "P50": _round(_pct(arr, 50), 6),
                "P75": _round(_pct(arr, 75), 6),
                "P95": _round(_pct(arr, 95), 6),
                "P99": _round(_pct(arr, 99), 6),
                "min": _round(float(arr.min()), 6),
                "max": _round(float(arr.max()), 6),
            }
        )
        recs = cat_ext.get(cat, [])
        if not recs:
            continue
        min_p = min(r[0] for r in recs)
        max_p = max(r[0] for r in recs)
        mins = sorted([r for r in recs if r[0] == min_p], key=lambda x: (x[1], x[2]))[:3]
        maxs = sorted([r for r in recs if r[0] == max_p], key=lambda x: (x[1], x[2]))[:3]
        for p, d, sku in mins:
            extra_rows.append(
                {
                    "品类": cat,
                    "极值类型": "min",
                    "单品编码": sku,
                    "单品名称": sku_to_name.get(sku, ""),
                    "日期": d.isoformat(),
                    "批发价格(元/千克)": _round(p, 6),
                }
            )
        for p, d, sku in maxs:
            extra_rows.append(
                {
                    "品类": cat,
                    "极值类型": "max",
                    "单品编码": sku,
                    "单品名称": sku_to_name.get(sku, ""),
                    "日期": d.isoformat(),
                    "批发价格(元/千克)": _round(p, 6),
                }
            )
    out = OUT / "04-品类已售批发价分位.csv"
    pd.DataFrame(rows).to_csv(out, index=False, encoding="utf-8-sig")
    out_ext = OUT / "04-极端批发价.csv"
    pd.DataFrame(extra_rows).to_csv(out_ext, index=False, encoding="utf-8-sig")
    print(f"  批发价对齐 已售命中={n_hit:,} 缺配={n_miss:,}", flush=True)
    return out, out_ext, series


def write_05(
    q3_sku_qty,
    q3_sku_days,
    wh,
    sku_to_cat,
    sku_to_name,
) -> Path:
    rows = []
    for sku in sorted(q3_sku_days, key=lambda s: (-len(q3_sku_days[s]), s)):
        days = q3_sku_days[sku]
        miss_days = sorted(d for d in days if (d, sku) not in wh)
        rows.append(
            {
                "单品编码": sku,
                "单品名称": sku_to_name.get(sku, ""),
                "分类名称": sku_to_cat.get(sku, "未知"),
                "窗口7天净销量_kg": _round(float(q3_sku_qty.get(sku, 0.0)), 6),
                "销售天数": len(days),
                "是否当天都有批发价": "是" if not miss_days else "否",
                "缺配天数": len(miss_days),
            }
        )
    out = OUT / "05-Q3窗口可售SKU.csv"
    pd.DataFrame(rows).to_csv(out, index=False, encoding="utf-8-sig")
    return out


def write_06(daily_cat_qty, cats: list[str]) -> tuple[Path, dict]:
    win_days = list(daterange(Q3_START, Q3_END))
    windows = {
        "窗口7": (Q3_START, Q3_END),
        "近30": (WIN_END - timedelta(days=29), WIN_END),
        "近90": (WIN_END - timedelta(days=89), WIN_END),
        "近365": (WIN_END - timedelta(days=364), WIN_END),
    }
    rows = []
    summary = {}
    for cat in cats:
        row = {"品类": cat}
        win_vals = []
        for d in win_days:
            v = float(daily_cat_qty.get(d, {}).get(cat, 0.0))
            row[d.isoformat()] = _round(v, 6)
            win_vals.append(v)
        win_arr = np.asarray(win_vals, dtype=float)
        row["窗口日均_kg"] = _round(_mean(win_arr), 6)
        row["窗口日标准差_kg"] = _round(_std_sample(win_arr), 6)
        cat_sum = {}
        for name, (a, b) in windows.items():
            arr = np.asarray(
                [float(daily_cat_qty.get(d, {}).get(cat, 0.0)) for d in daterange(a, b)],
                dtype=float,
            )
            row[f"{name}日均_kg"] = _round(_mean(arr), 6)
            row[f"{name}日标准差_kg"] = _round(_std_sample(arr), 6)
            row[f"{name}日历天数"] = int(arr.size)
            cat_sum[name] = {"mean": float(np.mean(arr)), "std": _std_sample(arr), "n": int(arr.size)}
        rows.append(row)
        summary[cat] = cat_sum
    # 窗口7 与 窗口日均 重复，保留窗口日均 + 近窗列；去掉「窗口7日均」以免双份
    df = pd.DataFrame(rows)
    if "窗口7日均_kg" in df.columns:
        df = df.drop(columns=["窗口7日均_kg", "窗口7日标准差_kg", "窗口7日历天数"])
    out = OUT / "06-窗口vs近期基线.csv"
    df.to_csv(out, index=False, encoding="utf-8-sig")
    return out, summary


def write_07(
    ret_cat_rows,
    ret_cat_qty,
    sale_cat_qty,
    ret_sku_rows,
    ret_sku_qty,
    sku_to_cat,
    sku_to_name,
    cats: list[str],
) -> tuple[Path, Path]:
    rows = []
    for cat in cats:
        n_ret = int(ret_cat_rows.get(cat, 0))
        q_ret = float(ret_cat_qty.get(cat, 0.0))
        q_sale = float(sale_cat_qty.get(cat, 0.0))
        ratio = abs(q_ret) / q_sale if q_sale else float("nan")
        rows.append(
            {
                "品类": cat,
                "退货行数": n_ret,
                "退货销量合计_kg": _round(q_ret, 6),
                "退货销量绝对值_kg": _round(abs(q_ret), 6),
                "销售类型销售_销量合计_kg": _round(q_sale, 6),
                "退货绝对值占销售销量比例": None if math.isnan(ratio) else _round(ratio, 8),
            }
        )
    out_cat = OUT / "07-退货影响.csv"
    pd.DataFrame(rows).to_csv(out_cat, index=False, encoding="utf-8-sig")

    skus = set(ret_sku_rows) | set(ret_sku_qty)
    recs = []
    for sku in skus:
        recs.append(
            {
                "单品编码": sku,
                "单品名称": sku_to_name.get(sku, ""),
                "分类名称": sku_to_cat.get(sku, "未知"),
                "退货行数": int(ret_sku_rows.get(sku, 0)),
                "退货销量合计_kg": _round(float(ret_sku_qty.get(sku, 0.0)), 6),
                "退货销量绝对值_kg": _round(abs(float(ret_sku_qty.get(sku, 0.0))), 6),
            }
        )
    recs.sort(key=lambda r: (-r["退货行数"], -r["退货销量绝对值_kg"], r["单品编码"]))
    for i, r in enumerate(recs, 1):
        r["行数排名"] = i
    by_abs = sorted(recs, key=lambda r: (-r["退货销量绝对值_kg"], -r["退货行数"], r["单品编码"]))
    abs_rank = {r["单品编码"]: i for i, r in enumerate(by_abs, 1)}
    for r in recs:
        r["绝对值销量排名"] = abs_rank[r["单品编码"]]
    top_row = {r["单品编码"] for r in recs if r["行数排名"] <= 10}
    top_abs = {r["单品编码"] for r in recs if r["绝对值销量排名"] <= 10}
    keep = [r for r in recs if r["单品编码"] in (top_row | top_abs)]
    keep.sort(key=lambda r: (r["行数排名"], r["绝对值销量排名"]))
    out_sku = OUT / "07-退货前10SKU.csv"
    pd.DataFrame(keep).to_csv(out_sku, index=False, encoding="utf-8-sig")
    return out_cat, out_sku


def write_08(df4: pd.DataFrame, sku_to_cat: dict[str, str], sku_to_name: dict[str, str], cats: list[str]) -> Path:
    df = df4.copy()
    df["品类"] = df["单品编码"].map(lambda x: sku_to_cat.get(x, "未知"))
    rows = []
    for cat in cats:
        sub = df[df["品类"] == cat]
        rates = sub["损耗率(%)"].dropna().to_numpy(dtype=float)
        n = int(rates.size)
        n_zero = int(np.sum(np.abs(rates) < 1e-12))
        n_943 = int(np.sum((np.round(rates, 2) == 9.43) | (np.abs(rates - 9.43) < 1e-6)))
        if n == 0:
            rows.append({"品类": cat, "n": 0})
            continue
        imax = int(np.argmax(rates))
        max_v = float(rates[imax])
        # 并列最大值
        names = []
        for sku, name, v in zip(sub["单品编码"], sub["单品名称"], sub["损耗率(%)"]):
            try:
                fv = float(v)
            except (TypeError, ValueError):
                continue
            if abs(fv - max_v) < 1e-9:
                names.append(str(name if pd.notna(name) else sku_to_name.get(sku, sku)))
        rows.append(
            {
                "品类": cat,
                "n": n,
                "P25": _round(_pct(rates, 25), 6),
                "P50": _round(_pct(rates, 50), 6),
                "P75": _round(_pct(rates, 75), 6),
                "P90": _round(_pct(rates, 90), 6),
                "损耗为0的个数": n_zero,
                "损耗恰好9.43的个数": n_943,
                "最大值": _round(max_v, 6),
                "最大值对应单品名": "；".join(names),
            }
        )
    out = OUT / "08-品类单品损耗分位.csv"
    pd.DataFrame(rows).to_csv(out, index=False, encoding="utf-8-sig")
    return out


def make_pngs(series02: dict[str, np.ndarray], summary06: dict, series04: dict[str, np.ndarray], cats: list[str]) -> list[Path]:
    font = _setup_cn_font()
    print(f"中文字体: {font}", flush=True)
    paths = []

    fig, ax = plt.subplots(figsize=(10, 5.5))
    data = [series02[c] for c in cats]
    bp = ax.boxplot(data, tick_labels=cats, showfliers=True, flierprops={"markersize": 3, "alpha": 0.4})
    ax.set_ylabel("日净销量 (kg)")
    ax.set_title("品类日净销量分布（2020-07-01～2023-06-30 日历日补0，非散点）")
    ax.grid(True, axis="y", alpha=0.25)
    fig.tight_layout()
    p = OUT / "02-品类日销量箱线.png"
    fig.savefig(p, dpi=120, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    paths.append(p)

    fig, ax = plt.subplots(figsize=(10, 5.5))
    x = np.arange(len(cats))
    w = 0.2
    m_win = [summary06[c]["窗口7"]["mean"] for c in cats]
    m_30 = [summary06[c]["近30"]["mean"] for c in cats]
    m_90 = [summary06[c]["近90"]["mean"] for c in cats]
    m_365 = [summary06[c]["近365"]["mean"] for c in cats]
    ax.bar(x - 1.5 * w, m_win, w, label="窗口 6/24–6/30 日均")
    ax.bar(x - 0.5 * w, m_30, w, label="近30日历日日均")
    ax.bar(x + 0.5 * w, m_90, w, label="近90日历日日均")
    ax.bar(x + 1.5 * w, m_365, w, label="近365日历日日均")
    ax.set_xticks(x)
    ax.set_xticklabels(cats)
    ax.set_ylabel("日均净销量 (kg)")
    ax.set_title("Q3 窗口日均 vs 截止 2023-06-30 的近 30/90/365 日均（日历日补0）")
    ax.legend(fontsize=8)
    ax.grid(True, axis="y", alpha=0.25)
    fig.tight_layout()
    p = OUT / "06-窗口vs365日均.png"
    fig.savefig(p, dpi=120, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    paths.append(p)

    fig, ax = plt.subplots(figsize=(10, 5.5))
    data = [series04.get(c, np.asarray([])) for c in cats]
    ax.boxplot(data, tick_labels=cats, showfliers=True, flierprops={"markersize": 3, "alpha": 0.35})
    ax.set_yscale("log")
    ax.set_ylabel("批发价 (元/kg，对数轴)")
    ax.set_title("已售 (日,SKU) 对齐附件3 的批发价分布（销售类型=销售）")
    ax.grid(True, axis="y", alpha=0.25, which="both")
    fig.tight_layout()
    p = OUT / "04-已售批发价箱线.png"
    fig.savefig(p, dpi=120, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    paths.append(p)
    return paths


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    print("读取附件1/3/4…", flush=True)
    df1, sku_to_cat, sku_to_name, cat_n_sku = load_att1()
    cats = [c for c in CAT_ORDER if c in set(df1["分类名称"].astype(str))]
    extra = [c for c in df1["分类名称"].astype(str).unique() if c not in cats]
    cats = cats + extra
    wh = load_att3()
    df4 = load_att4_sku()
    print(f"  附件1 SKU={len(sku_to_cat)} 品类={cats}", flush=True)
    print(f"  附件3 (日,SKU)={len(wh):,}", flush=True)
    print(f"  附件4 单品={len(df4)}", flush=True)

    print("流式扫描附件2…", flush=True)
    agg = stream_att2(sku_to_cat)

    print("写 CSV…", flush=True)
    p01 = write_01(agg["daily_cat_qty"], agg["cat_sale_days"], cats)
    p02, series02 = write_02(agg["daily_cat_qty"], agg["cat_skus_flow"], cat_n_sku, cats)
    p03 = write_03(agg["pearson"], cats)
    p04, p04e, series04 = write_04(agg["sold_pairs"], wh, sku_to_cat, sku_to_name, cats)
    p05 = write_05(agg["q3_sku_qty"], agg["q3_sku_days"], wh, sku_to_cat, sku_to_name)
    p06, summary06 = write_06(agg["daily_cat_qty"], cats)
    p07, p07s = write_07(
        agg["ret_cat_rows"],
        agg["ret_cat_qty"],
        agg["sale_cat_qty"],
        agg["ret_sku_rows"],
        agg["ret_sku_qty"],
        sku_to_cat,
        sku_to_name,
        cats,
    )
    p08 = write_08(df4, sku_to_cat, sku_to_name, cats)
    pngs = make_pngs(series02, summary06, series04, cats)

    print("=== 产出 ===", flush=True)
    for p in [p01, p02, p03, p04, p04e, p05, p06, p07, p07s, p08, *pngs]:
        print(f"  {p.name}  {p.stat().st_size:,} bytes", flush=True)
    print(f"Q3 窗口 SKU 数={len(agg['q3_sku_days'])}", flush=True)
    print(f"已售 (日,SKU)={len(agg['sold_pairs'])}", flush=True)


if __name__ == "__main__":
    main()
