"""索取第 2 轮：附件2 只 openpyxl read_only 流式扫描一次；不覆盖 01–08。"""
from __future__ import annotations

import math
import sys
from collections import defaultdict
from datetime import date, timedelta
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.stats import spearmanr

_HERE = Path(__file__).resolve().parent
if str(_HERE) not in sys.path:
    sys.path.insert(0, str(_HERE))

from _compute import (  # noqa: E402
    CAT_ORDER,
    OUT,
    SOURCE,
    _as_date,
    _mean,
    _norm_code,
    _pct,
    _round,
    _setup_cn_font,
    _std_sample,
    daterange,
    load_att1,
    load_att3,
)

JUN_START = date(2023, 6, 1)
JUN_END = date(2023, 6, 30)
Q3_START = date(2023, 6, 24)
Q3_END = date(2023, 6, 30)
NEAR30_START = date(2023, 6, 1)  # 截止 06-30 含当天 30 个日历日
NEAR90_START = date(2023, 4, 2)  # 截止 06-30 含当天 90 个日历日
R13_START = date(2023, 4, 1)
R15_START = date(2023, 4, 1)
SCAN_START = date(2023, 4, 1)
WIN_END = date(2023, 6, 30)

WD_CN = ["周一", "周二", "周三", "周四", "周五", "周六", "周日"]
JUN_DAYS = list(daterange(JUN_START, JUN_END))
WIN_DAYS = list(daterange(Q3_START, Q3_END))
NEAR30_DAYS = list(daterange(NEAR30_START, WIN_END))
NEAR90_DAYS = list(daterange(NEAR90_START, WIN_END))
R13_DAYS = list(daterange(R13_START, WIN_END))
R15_DAYS = list(daterange(R15_START, WIN_END))


def _is_zero(x: float) -> bool:
    return abs(float(x)) < 1e-12


def load_q3_skus() -> list[str]:
    p = OUT / "05-Q3窗口可售SKU.csv"
    df = pd.read_csv(p, dtype={"单品编码": str})
    df["单品编码"] = df["单品编码"].map(_norm_code)
    skus = [s for s in df["单品编码"].tolist() if s]
    if len(skus) != 49:
        print(f"  警告: 05 文件 SKU 数={len(skus)}，预期 49", flush=True)
    return skus


def stream_att2(sku_to_cat: dict[str, str]) -> dict:
    p = SOURCE / "附件2.xlsx"
    wb = load_workbook(p, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    header = [h if h is None else str(h).strip() for h in next(it)]
    idx = {name: i for i, name in enumerate(header)}
    required = ["销售日期", "单品编码", "销量(千克)", "销售单价(元/千克)", "销售类型"]
    missing_cols = [c for c in required if c not in idx]
    if missing_cols:
        wb.close()
        raise KeyError(f"附件2 缺列: {missing_cols}; header={header}")

    i_date = idx["销售日期"]
    i_sku = idx["单品编码"]
    i_qty = idx["销量(千克)"]
    i_price = idx["销售单价(元/千克)"]
    i_type = idx["销售类型"]

    sku_day_qty: dict[str, dict[date, float]] = defaultdict(lambda: defaultdict(float))
    sku_day_sale_qty: dict[str, dict[date, float]] = defaultdict(lambda: defaultdict(float))
    sku_day_sale_amt: dict[str, dict[date, float]] = defaultdict(lambda: defaultdict(float))
    sku_june_days: dict[str, set[date]] = defaultdict(set)
    sku_win_days: dict[str, set[date]] = defaultdict(set)
    daily_cat_qty: dict[date, dict[str, float]] = defaultdict(lambda: defaultdict(float))

    n = 0
    n_used = 0
    for raw in it:
        if raw is None or all(v is None for v in raw):
            continue
        n += 1
        d = _as_date(raw[i_date] if i_date < len(raw) else None)
        if d is None or d < SCAN_START or d > WIN_END:
            if n % 200_000 == 0:
                print(f"  附件2 已扫描 {n:,} 行…", flush=True)
            continue
        sku = _norm_code(raw[i_sku] if i_sku < len(raw) else None)
        if not sku:
            continue
        stype = str(raw[i_type]).strip() if raw[i_type] is not None else ""
        try:
            qv = float(raw[i_qty]) if raw[i_qty] is not None else None
        except (TypeError, ValueError):
            qv = None
        try:
            pv = float(raw[i_price]) if raw[i_price] is not None else None
        except (TypeError, ValueError):
            pv = None

        n_used += 1
        cat = sku_to_cat.get(sku, "未知")

        if qv is not None:
            sku_day_qty[sku][d] += qv
            daily_cat_qty[d][cat] += qv

        if JUN_START <= d <= JUN_END:
            sku_june_days[sku].add(d)
        if Q3_START <= d <= Q3_END:
            sku_win_days[sku].add(d)

        if stype == "销售" and qv is not None and qv > 0 and pv is not None and not math.isnan(pv):
            sku_day_sale_qty[sku][d] += qv
            sku_day_sale_amt[sku][d] += qv * pv

        if n % 200_000 == 0:
            print(f"  附件2 已扫描 {n:,} 行…", flush=True)

    wb.close()
    print(f"  附件2 扫描完成 {n:,} 行；2023-04-01～06-30 有效行={n_used:,}", flush=True)
    return {
        "n": n,
        "n_used": n_used,
        "sku_day_qty": sku_day_qty,
        "sku_day_sale_qty": sku_day_sale_qty,
        "sku_day_sale_amt": sku_day_sale_amt,
        "sku_june_days": sku_june_days,
        "sku_win_days": sku_win_days,
        "daily_cat_qty": daily_cat_qty,
    }


def _sku_sum(sku_day_qty, sku: str, days: list[date]) -> float:
    m = sku_day_qty.get(sku, {})
    return float(sum(m.get(d, 0.0) for d in days))


def _sku_arr(sku_day_qty, sku: str, days: list[date]) -> np.ndarray:
    m = sku_day_qty.get(sku, {})
    return np.asarray([float(m.get(d, 0.0)) for d in days], dtype=float)


def write_11(agg, sku_to_cat, sku_to_name, cats: list[str]) -> tuple[Path, Path, dict]:
    sku_day_qty = agg["sku_day_qty"]
    sku_june_days = agg["sku_june_days"]
    sku_win_days = agg["sku_win_days"]
    daily_cat_qty = agg["daily_cat_qty"]

    skus = sorted(set(sku_june_days) | set(sku_win_days))
    rows = []
    for sku in skus:
        june_days_n = len(sku_june_days.get(sku, set()))
        win_days_n = len(sku_win_days.get(sku, set()))
        if june_days_n == 0 and win_days_n == 0:
            continue
        arr = _sku_arr(sku_day_qty, sku, JUN_DAYS)
        win_qty = _sku_sum(sku_day_qty, sku, WIN_DAYS)
        rows.append(
            {
                "单品编码": sku,
                "单品名称": sku_to_name.get(sku, ""),
                "分类名称": sku_to_cat.get(sku, "未知"),
                "六月日均净销量_日历30补0_kg": _round(_mean(arr), 6),
                "六月销售天数": june_days_n,
                "六月日销量标准差_日历30补0_kg": _round(_std_sample(arr), 6),
                "窗口7日净销量合计_kg": _round(win_qty, 6),
                "窗口销售天数": win_days_n,
            }
        )
    rows.sort(key=lambda r: (cats.index(r["分类名称"]) if r["分类名称"] in cats else 99, r["单品编码"]))
    out_sku = OUT / "11-六月SKU与窗口.csv"
    pd.DataFrame(rows).to_csv(out_sku, index=False, encoding="utf-8-sig")

    cat_rows = []
    summary = {}
    for cat in cats:
        june_arr = np.asarray(
            [float(daily_cat_qty.get(d, {}).get(cat, 0.0)) for d in JUN_DAYS], dtype=float
        )
        win_total = float(
            sum(float(daily_cat_qty.get(d, {}).get(cat, 0.0)) for d in WIN_DAYS)
        )
        june_sold = [
            r for r in rows if r["分类名称"] == cat and r["六月销售天数"] > 0
        ]
        win_sold = [
            r for r in rows if r["分类名称"] == cat and r["窗口销售天数"] > 0
        ]
        drop = [
            r
            for r in june_sold
            if r["窗口7日净销量合计_kg"] is None or _is_zero(r["窗口7日净销量合计_kg"])
        ]
        n_june = len(june_sold)
        n_drop = len(drop)
        ratio = n_drop / n_june if n_june else float("nan")
        cat_rows.append(
            {
                "品类": cat,
                "六月日均_kg": _round(_mean(june_arr), 6),
                "窗口7日合计_kg": _round(win_total, 6),
                "窗口有销售SKU数": len(win_sold),
                "六月有售SKU数": n_june,
                "六月有售但窗口销量为0_SKU数": n_drop,
                "掉SKU占六月有售比例": None if math.isnan(ratio) else _round(ratio, 6),
            }
        )
        summary[cat] = {
            "june_mean": _mean(june_arr),
            "win_total": win_total,
            "n_win": len(win_sold),
            "n_june": n_june,
            "n_drop": n_drop,
            "drop_ratio": ratio,
        }
    out_cat = OUT / "11-品类构成分解.csv"
    pd.DataFrame(cat_rows).to_csv(out_cat, index=False, encoding="utf-8-sig")
    return out_sku, out_cat, summary


def write_12(agg, q3_skus, sku_to_cat, sku_to_name, cats: list[str]) -> tuple[Path, Path, dict]:
    sku_day_qty = agg["sku_day_qty"]
    n30 = len(NEAR30_DAYS)
    n90 = len(NEAR90_DAYS)
    n7 = len(WIN_DAYS)
    rows = []
    for sku in q3_skus:
        win_total = _sku_sum(sku_day_qty, sku, WIN_DAYS)
        win_mean = win_total / n7 if n7 else float("nan")
        m30 = _mean(_sku_arr(sku_day_qty, sku, NEAR30_DAYS))
        m90 = _mean(_sku_arr(sku_day_qty, sku, NEAR90_DAYS))
        r30 = win_mean / m30 if m30 and not _is_zero(m30) and not math.isnan(m30) else float("nan")
        r90 = win_mean / m90 if m90 and not _is_zero(m90) and not math.isnan(m90) else float("nan")
        rows.append(
            {
                "单品编码": sku,
                "单品名称": sku_to_name.get(sku, ""),
                "分类名称": sku_to_cat.get(sku, "未知"),
                "窗口7日净销量合计_kg": _round(win_total, 6),
                "窗口日均_日历7补0_kg": _round(win_mean, 6),
                "近30日均_日历补0_kg": _round(m30, 6),
                "近90日均_日历补0_kg": _round(m90, 6),
                "比值_窗口日均_近30": None if math.isnan(r30) else _round(r30, 6),
                "比值_窗口日均_近90": None if math.isnan(r90) else _round(r90, 6),
            }
        )
    out_sku = OUT / "12-窗口SKU相对自身.csv"
    pd.DataFrame(rows).to_csv(out_sku, index=False, encoding="utf-8-sig")

    cat_rows = []
    summary = {}
    for cat in cats:
        sub = [r for r in rows if r["分类名称"] == cat]
        a30 = np.asarray(
            [float(r["比值_窗口日均_近30"]) for r in sub if r["比值_窗口日均_近30"] is not None],
            dtype=float,
        )
        a90 = np.asarray(
            [float(r["比值_窗口日均_近90"]) for r in sub if r["比值_窗口日均_近90"] is not None],
            dtype=float,
        )
        cat_rows.append(
            {
                "品类": cat,
                "n_窗口SKU": len(sub),
                "n_比值近30": int(a30.size),
                "比值近30_P25": _round(_pct(a30, 25), 6) if a30.size else None,
                "比值近30_P50": _round(_pct(a30, 50), 6) if a30.size else None,
                "比值近30_P75": _round(_pct(a30, 75), 6) if a30.size else None,
                "n_比值近90": int(a90.size),
                "比值近90_P25": _round(_pct(a90, 25), 6) if a90.size else None,
                "比值近90_P50": _round(_pct(a90, 50), 6) if a90.size else None,
                "比值近90_P75": _round(_pct(a90, 75), 6) if a90.size else None,
            }
        )
        summary[cat] = {
            "n": len(sub),
            "r30_p50": _pct(a30, 50) if a30.size else float("nan"),
            "r90_p50": _pct(a90, 50) if a90.size else float("nan"),
            "a30": a30,
            "a90": a90,
        }
    out_cat = OUT / "12-品类比值分位.csv"
    pd.DataFrame(cat_rows).to_csv(out_cat, index=False, encoding="utf-8-sig")
    return out_sku, out_cat, summary


def _corr_pair(qty: np.ndarray, price: np.ndarray) -> tuple[float, float]:
    if qty.size < 10:
        return float("nan"), float("nan")
    if float(np.std(qty, ddof=0)) <= 0 or float(np.std(price, ddof=0)) <= 0:
        return float("nan"), float("nan")
    pr = float(np.corrcoef(qty, price)[0, 1])
    sr = float(spearmanr(qty, price).correlation)
    return pr, sr


def write_13(agg, sku_to_cat, sku_to_name, cats: list[str]) -> tuple[Path, Path, dict]:
    sku_day_qty = agg["sku_day_qty"]
    sku_day_sale_qty = agg["sku_day_sale_qty"]
    sku_day_sale_amt = agg["sku_day_sale_amt"]

    recs = []
    for sku, day_map in sku_day_qty.items():
        xs = []
        ys = []
        for d in R13_DAYS:
            sq = float(sku_day_sale_qty.get(sku, {}).get(d, 0.0))
            if sq <= 0:
                continue
            amt = float(sku_day_sale_amt.get(sku, {}).get(d, 0.0))
            price = amt / sq
            qty = float(day_map.get(d, 0.0))
            xs.append(qty)
            ys.append(price)
        n = len(xs)
        if n < 10:
            continue
        pr, sr = _corr_pair(np.asarray(xs, dtype=float), np.asarray(ys, dtype=float))
        recs.append(
            {
                "单品编码": sku,
                "单品名称": sku_to_name.get(sku, ""),
                "分类名称": sku_to_cat.get(sku, "未知"),
                "n_日": n,
                "pearson": pr,
                "spearman": sr,
            }
        )

    cat_rows = []
    summary = {}
    for cat in cats:
        sub = [r for r in recs if r["分类名称"] == cat]
        pr = np.asarray([r["pearson"] for r in sub if not math.isnan(r["pearson"])], dtype=float)
        sr = np.asarray([r["spearman"] for r in sub if not math.isnan(r["spearman"])], dtype=float)
        n_strong = int(np.sum(np.abs(pr) > 0.3)) if pr.size else 0
        cat_rows.append(
            {
                "品类": cat,
                "参与SKU数_n>=10": len(sub),
                "n_Pearson有效": int(pr.size),
                "Pearson_P25": _round(_pct(pr, 25), 6) if pr.size else None,
                "Pearson_P50": _round(_pct(pr, 50), 6) if pr.size else None,
                "Pearson_P75": _round(_pct(pr, 75), 6) if pr.size else None,
                "n_Spearman有效": int(sr.size),
                "Spearman_P25": _round(_pct(sr, 25), 6) if sr.size else None,
                "Spearman_P50": _round(_pct(sr, 50), 6) if sr.size else None,
                "Spearman_P75": _round(_pct(sr, 75), 6) if sr.size else None,
                "absPearson大于0.3个数": n_strong,
            }
        )
        summary[cat] = {
            "n": len(sub),
            "p50": _pct(pr, 50) if pr.size else float("nan"),
            "n_strong": n_strong,
        }
    out_cat = OUT / "13-同SKU量价-品类分布.csv"
    pd.DataFrame(cat_rows).to_csv(out_cat, index=False, encoding="utf-8-sig")

    valid = [r for r in recs if not math.isnan(r["pearson"])]
    pos = sorted([r for r in valid if r["pearson"] > 0], key=lambda r: -r["pearson"])[:3]
    neg = sorted([r for r in valid if r["pearson"] < 0], key=lambda r: r["pearson"])[:3]
    ext_rows = []
    for tag, group in (("正", pos), ("负", neg)):
        for r in group:
            ext_rows.append(
                {
                    "极值类型": tag,
                    "单品编码": r["单品编码"],
                    "单品名称": r["单品名称"],
                    "分类名称": r["分类名称"],
                    "n_日": r["n_日"],
                    "pearson": _round(r["pearson"], 6),
                    "spearman": None if math.isnan(r["spearman"]) else _round(r["spearman"], 6),
                }
            )
    out_ext = OUT / "13-同SKU量价-极端例.csv"
    pd.DataFrame(ext_rows).to_csv(out_ext, index=False, encoding="utf-8-sig")
    summary["_n_total"] = len(recs)
    summary["_ext"] = ext_rows
    return out_cat, out_ext, summary


def write_14(q3_skus, wh, sku_to_cat, sku_to_name, cats: list[str]) -> tuple[Path, Path, dict]:
    rows = []
    n_miss = 0
    n_hit = 0
    for sku in q3_skus:
        prices = []
        miss_days = []
        for d in WIN_DAYS:
            p = wh.get((d, sku))
            if p is None or (isinstance(p, float) and math.isnan(p)):
                miss_days.append(d)
                n_miss += 1
                continue
            prices.append(float(p))
            n_hit += 1
        arr = np.asarray(prices, dtype=float)
        mean_p = _mean(arr) if arr.size else float("nan")
        std_p = _std_sample(arr) if arr.size else float("nan")
        if arr.size < 2 or math.isnan(mean_p) or _is_zero(mean_p) or math.isnan(std_p):
            cv = float("nan")
        else:
            cv = std_p / mean_p
        rows.append(
            {
                "单品编码": sku,
                "单品名称": sku_to_name.get(sku, ""),
                "分类名称": sku_to_cat.get(sku, "未知"),
                "n_有批发价天数": int(arr.size),
                "缺配天数": len(miss_days),
                "批发价_P25": _round(_pct(arr, 25), 6) if arr.size else None,
                "批发价_P50": _round(_pct(arr, 50), 6) if arr.size else None,
                "批发价_P75": _round(_pct(arr, 75), 6) if arr.size else None,
                "批发价_7日均值": _round(mean_p, 6) if arr.size else None,
                "CV_std_over_mean": None if math.isnan(cv) else _round(cv, 6),
            }
        )
    out_sku = OUT / "14-窗口SKU批发价.csv"
    pd.DataFrame(rows).to_csv(out_sku, index=False, encoding="utf-8-sig")
    print(f"  窗口7日×49 SKU 批发价 命中={n_hit} 缺配={n_miss}", flush=True)

    cat_rows = []
    summary = {}
    for cat in cats:
        sub = [r for r in rows if r["分类名称"] == cat]
        cvs = np.asarray(
            [float(r["CV_std_over_mean"]) for r in sub if r["CV_std_over_mean"] is not None],
            dtype=float,
        )
        p50s = np.asarray(
            [float(r["批发价_P50"]) for r in sub if r["批发价_P50"] is not None],
            dtype=float,
        )
        cat_rows.append(
            {
                "品类": cat,
                "SKU数": len(sub),
                "n_有CV": int(cvs.size),
                "CV_P50": _round(_pct(cvs, 50), 6) if cvs.size else None,
                "CV_P75": _round(_pct(cvs, 75), 6) if cvs.size else None,
                "批发价P50的中位": _round(_pct(p50s, 50), 6) if p50s.size else None,
            }
        )
        summary[cat] = {
            "n": len(sub),
            "cv_p50": _pct(cvs, 50) if cvs.size else float("nan"),
            "cv_p75": _pct(cvs, 75) if cvs.size else float("nan"),
            "p50_med": _pct(p50s, 50) if p50s.size else float("nan"),
        }
    out_cat = OUT / "14-品类成本稳定.csv"
    pd.DataFrame(cat_rows).to_csv(out_cat, index=False, encoding="utf-8-sig")
    return out_sku, out_cat, summary


def write_15(agg, cats: list[str]) -> tuple[Path, Path, dict]:
    daily_cat_qty = agg["daily_cat_qty"]
    n_cal = len(R15_DAYS)
    by_wd: dict[str, dict[int, list[float]]] = {c: {i: [] for i in range(7)} for c in cats}
    overall = {}
    for cat in cats:
        arr = np.asarray(
            [float(daily_cat_qty.get(d, {}).get(cat, 0.0)) for d in R15_DAYS], dtype=float
        )
        overall[cat] = _mean(arr)
        for d in R15_DAYS:
            by_wd[cat][d.weekday()].append(float(daily_cat_qty.get(d, {}).get(cat, 0.0)))

    rows = []
    for cat in cats:
        for wd in range(7):
            arr = np.asarray(by_wd[cat][wd], dtype=float)
            rows.append(
                {
                    "品类": cat,
                    "星期几_code": wd,
                    "星期几": WD_CN[wd],
                    "天数": int(arr.size),
                    "日均净销量_kg": _round(_mean(arr), 6),
                    "日销量标准差_kg": _round(_std_sample(arr), 6),
                    "同期总体日均_kg": _round(overall[cat], 6),
                    "相对总体": _round(_mean(arr) / overall[cat], 6)
                    if overall[cat] and not _is_zero(overall[cat])
                    else None,
                }
            )
    out_csv = OUT / "15-星期效应.csv"
    pd.DataFrame(rows).to_csv(out_csv, index=False, encoding="utf-8-sig")

    font = _setup_cn_font()
    print(f"  中文字体: {font}", flush=True)
    mat = np.zeros((len(cats), 7), dtype=float)
    for i, cat in enumerate(cats):
        for wd in range(7):
            mat[i, wd] = _mean(np.asarray(by_wd[cat][wd], dtype=float))
    fig, ax = plt.subplots(figsize=(10, 4.8))
    im = ax.imshow(mat, aspect="auto", cmap="YlOrRd")
    ax.set_xticks(range(7))
    ax.set_xticklabels(WD_CN)
    ax.set_yticks(range(len(cats)))
    ax.set_yticklabels(cats)
    ax.set_title("2023-04-01～06-30 品类×星期几 日均净销量（日历补0，kg）")
    for i in range(len(cats)):
        for j in range(7):
            ax.text(j, i, f"{mat[i, j]:.1f}", ha="center", va="center", fontsize=8, color="black")
    fig.colorbar(im, ax=ax, fraction=0.03, pad=0.02, label="日均 kg")
    fig.tight_layout()
    out_png = OUT / "15-星期效应.png"
    fig.savefig(out_png, dpi=120, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return out_csv, out_png, {"overall": overall, "n_cal": n_cal}


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    print("读取附件1/3 与 05 窗口名单…", flush=True)
    df1, sku_to_cat, sku_to_name, _cat_n = load_att1()
    cats = [c for c in CAT_ORDER if c in set(df1["分类名称"].astype(str))]
    extra = [c for c in df1["分类名称"].astype(str).unique() if c not in cats]
    cats = cats + extra
    wh = load_att3()
    q3_skus = load_q3_skus()
    print(f"  附件1 SKU={len(sku_to_cat)} 品类={cats}", flush=True)
    print(f"  附件3 (日,SKU)={len(wh):,}", flush=True)
    print(f"  Q3 窗口名单={len(q3_skus)}", flush=True)
    print(f"  近30={NEAR30_DAYS[0]}～{NEAR30_DAYS[-1]} n={len(NEAR30_DAYS)}", flush=True)
    print(f"  近90={NEAR90_DAYS[0]}～{NEAR90_DAYS[-1]} n={len(NEAR90_DAYS)}", flush=True)
    print(f"  条13/15={R15_DAYS[0]}～{R15_DAYS[-1]} n={len(R15_DAYS)}", flush=True)

    print("流式扫描附件2（仅聚合 2023-04-01～06-30）…", flush=True)
    agg = stream_att2(sku_to_cat)

    print("写 11–15…", flush=True)
    p11a, p11b, s11 = write_11(agg, sku_to_cat, sku_to_name, cats)
    p12a, p12b, s12 = write_12(agg, q3_skus, sku_to_cat, sku_to_name, cats)
    p13a, p13b, s13 = write_13(agg, sku_to_cat, sku_to_name, cats)
    p14a, p14b, s14 = write_14(q3_skus, wh, sku_to_cat, sku_to_name, cats)
    p15a, p15p, s15 = write_15(agg, cats)

    print("=== 产出 ===", flush=True)
    for p in [p11a, p11b, p12a, p12b, p13a, p13b, p14a, p14b, p15a, p15p]:
        print(f"  {p.name}  {p.stat().st_size:,} bytes", flush=True)

    print("=== 11 品类构成 ===", flush=True)
    for cat in cats:
        s = s11[cat]
        print(
            f"  {cat}: 六月日均={s['june_mean']:.3f} 窗口合计={s['win_total']:.3f} "
            f"窗口SKU={s['n_win']} 六月有售={s['n_june']} 掉SKU={s['n_drop']} "
            f"占比={s['drop_ratio']:.3f}",
            flush=True,
        )
    print("=== 12 比值中位 ===", flush=True)
    for cat in cats:
        s = s12[cat]
        print(
            f"  {cat}: n={s['n']} 近30中位={s['r30_p50']:.3f} 近90中位={s['r90_p50']:.3f}",
            flush=True,
        )
    print(f"=== 13 参与 SKU={s13['_n_total']} ===", flush=True)
    for cat in cats:
        s = s13[cat]
        print(f"  {cat}: n={s['n']} Pearson中位={s['p50']:.3f} |r|>0.3={s['n_strong']}", flush=True)
    print("=== 14 CV ===", flush=True)
    for cat in cats:
        s = s14[cat]
        print(
            f"  {cat}: n={s['n']} CV中位={s['cv_p50']:.4f} CV_P75={s['cv_p75']:.4f} "
            f"P50中位={s['p50_med']:.3f}",
            flush=True,
        )
    print(f"=== 15 91日总体日均 n={s15['n_cal']} ===", flush=True)
    for cat in cats:
        print(f"  {cat}: {s15['overall'][cat]:.3f}", flush=True)
    print("11 SKU 行数=", sum(1 for _ in open(p11a, encoding="utf-8-sig")) - 1, flush=True)


if __name__ == "__main__":
    main()
