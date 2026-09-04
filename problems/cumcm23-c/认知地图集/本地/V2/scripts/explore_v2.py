"""V2 本地探索：品类日相关、量-加成、历年7月初、Q3窗口约束。
只读题根 source/；写出 V2/data 与 V2/figures。
"""
from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib import font_manager
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[4]  # problems/cumcm23-c
SOURCE = ROOT / "source"
OUT_DATA = Path(__file__).resolve().parents[1] / "data"
OUT_FIG = Path(__file__).resolve().parents[1] / "figures"
OUT_DATA.mkdir(parents=True, exist_ok=True)
OUT_FIG.mkdir(parents=True, exist_ok=True)

CAT_ORDER = ["花叶类", "花菜类", "水生根茎类", "茄类", "辣椒类", "食用菌"]
Q3_START, Q3_END = date(2023, 6, 24), date(2023, 6, 30)


def setup_font() -> None:
    names = {f.name for f in font_manager.fontManager.ttflist}
    chosen = next((n for n in ("Microsoft YaHei", "SimHei") if n in names), "DejaVu Sans")
    plt.rcParams["font.sans-serif"] = [chosen, "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False


def norm_code(x) -> str | None:
    if x is None:
        return None
    if isinstance(x, float):
        if np.isnan(x):
            return None
        return str(int(x)) if x.is_integer() else str(x).strip()
    if isinstance(x, int):
        return str(x)
    s = str(x).strip()
    return s or None


def as_date(x) -> date | None:
    if x is None:
        return None
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    try:
        return pd.Timestamp(x).date()
    except Exception:
        return None


def load_catalog() -> pd.DataFrame:
    df = pd.read_excel(SOURCE / "附件1.xlsx", dtype=str)
    df.columns = ["单品编码", "单品名称", "分类编码", "分类名称"]
    df["单品编码"] = df["单品编码"].map(norm_code)
    return df


def load_loss() -> tuple[pd.DataFrame, pd.DataFrame]:
    cat = pd.read_excel(SOURCE / "附件4.xlsx", sheet_name="平均损耗率(%)_小分类编码_不同值")
    cat.columns = ["分类编码", "分类名称", "品类损耗率pct"]
    sku = pd.read_excel(SOURCE / "附件4.xlsx", sheet_name="Sheet1")
    sku.columns = ["单品编码", "单品名称", "单品损耗率pct"]
    sku["单品编码"] = sku["单品编码"].map(norm_code)
    return cat, sku


def stream_sales(sku2cat: dict[str, str]):
    """一次流式扫描附件2，累计品类日净销量、金额、批发对齐所需日SKU净销量、窗口SKU。"""
    cat_day_qty: dict[tuple[date, str], float] = defaultdict(float)
    cat_day_rev: dict[tuple[date, str], float] = defaultdict(float)
    cat_day_disc_qty: dict[tuple[date, str], float] = defaultdict(float)
    day_sku_qty: dict[tuple[date, str], float] = defaultdict(float)
    day_sku_rev: dict[tuple[date, str], float] = defaultdict(float)
    win_sku_qty: dict[str, float] = defaultdict(float)
    win_sku_days: dict[str, set] = defaultdict(set)
    july_hist: dict[tuple[int, date, str], float] = defaultdict(float)  # year, calendar-day-of-week mapped?

    # 历年 7/1–7/7
    july_targets = {
        y: {date(y, 7, d) for d in range(1, 8)} for y in (2020, 2021, 2022)
    }

    wb = load_workbook(SOURCE / "附件2.xlsx", read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    # 销售日期, 扫码销售时间, 单品编码, 销量(千克), 销售单价(元/千克), 销售类型, 是否打折销售
    n = 0
    for row in rows:
        n += 1
        d = as_date(row[0])
        code = norm_code(row[2])
        qty = float(row[3] or 0)
        price = float(row[4] or 0)
        disc = row[6]
        if d is None or code is None:
            continue
        cat = sku2cat.get(code)
        if cat is None:
            continue
        cat_day_qty[(d, cat)] += qty
        cat_day_rev[(d, cat)] += qty * price
        if disc == "是":
            cat_day_disc_qty[(d, cat)] += qty
        day_sku_qty[(d, code)] += qty
        day_sku_rev[(d, code)] += qty * price
        if Q3_START <= d <= Q3_END:
            win_sku_qty[code] += qty
            if qty != 0:
                win_sku_days[code].add(d)
        for y, days in july_targets.items():
            if d in days:
                july_hist[(y, d, cat)] += qty
    wb.close()
    return {
        "n_rows": n,
        "cat_day_qty": cat_day_qty,
        "cat_day_rev": cat_day_rev,
        "cat_day_disc_qty": cat_day_disc_qty,
        "day_sku_qty": day_sku_qty,
        "day_sku_rev": day_sku_rev,
        "win_sku_qty": win_sku_qty,
        "win_sku_days": win_sku_days,
        "july_hist": july_hist,
    }


def load_wholesale_for_days(days: set[date]) -> dict[tuple[date, str], float]:
    """只保留关心日期的批发价。"""
    out: dict[tuple[date, str], float] = {}
    wb = load_workbook(SOURCE / "附件3.xlsx", read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    next(rows)
    for row in rows:
        d = as_date(row[0])
        if d is None or d not in days:
            continue
        code = norm_code(row[1])
        if code is None:
            continue
        out[(d, code)] = float(row[2])
    wb.close()
    return out


def pearson(xs: np.ndarray, ys: np.ndarray) -> float:
    if xs.size < 5:
        return float("nan")
    if np.std(xs) == 0 or np.std(ys) == 0:
        return float("nan")
    return float(np.corrcoef(xs, ys)[0, 1])


def main() -> None:
    setup_font()
    catalog = load_catalog()
    sku2cat = dict(zip(catalog["单品编码"], catalog["分类名称"]))
    sku2name = dict(zip(catalog["单品编码"], catalog["单品名称"]))
    loss_cat, loss_sku = load_loss()

    print("streaming 附件2 …")
    S = stream_sales(sku2cat)
    print(f"rows={S['n_rows']}")

    # --- 品类日面板（日历补 0） ---
    all_days = pd.date_range("2020-07-01", "2023-06-30", freq="D").date
    records = []
    for d in all_days:
        for cat in CAT_ORDER:
            q = S["cat_day_qty"].get((d, cat), 0.0)
            r = S["cat_day_rev"].get((d, cat), 0.0)
            dq = S["cat_day_disc_qty"].get((d, cat), 0.0)
            records.append(
                {
                    "日期": d,
                    "品类": cat,
                    "净销量_kg": q,
                    "销售额_元": r,
                    "打折销量_kg": dq,
                    "均价_元每kg": (r / q) if q != 0 else np.nan,
                    "星期": d.weekday(),
                }
            )
    panel = pd.DataFrame(records)
    panel.to_csv(OUT_DATA / "品类日面板.csv", index=False, encoding="utf-8-sig")

    # 近 365 / 近 90 / 全期 品类间日销量相关
    def corr_block(start: date, end: date, tag: str) -> pd.DataFrame:
        sub = panel[(panel["日期"] >= start) & (panel["日期"] <= end)]
        wide = sub.pivot(index="日期", columns="品类", values="净销量_kg")[CAT_ORDER]
        corr = wide.corr()
        corr.to_csv(OUT_DATA / f"品类日相关_{tag}.csv", encoding="utf-8-sig")
        return corr

    c365 = corr_block(date(2022, 7, 1), date(2023, 6, 30), "近365")
    c90 = corr_block(date(2023, 4, 2), date(2023, 6, 30), "近90")
    c_all = corr_block(date(2020, 7, 1), date(2023, 6, 30), "全期")
    print("corr365 flower-leaf vs chili:", c365.loc["花叶类", "辣椒类"])
    print("corr365 leaf vs fungi:", c365.loc["花叶类", "食用菌"])

    # 图：近365相关热力
    fig, ax = plt.subplots(figsize=(7, 5.5))
    im = ax.imshow(c365.values, cmap="RdBu_r", vmin=-1, vmax=1)
    ax.set_xticks(range(6))
    ax.set_yticks(range(6))
    ax.set_xticklabels(CAT_ORDER, rotation=30, ha="right")
    ax.set_yticklabels(CAT_ORDER)
    for i in range(6):
        for j in range(6):
            ax.text(j, i, f"{c365.values[i, j]:.2f}", ha="center", va="center", fontsize=8)
    ax.set_title("品类日净销量相关（2022-07-01～2023-06-30）")
    fig.colorbar(im, ax=ax, fraction=0.046)
    fig.tight_layout()
    fig.savefig(OUT_FIG / "品类日相关_近365.png", dpi=120)
    plt.close(fig)

    # --- 量-加成：用近90日已售(日,SKU)对齐批发价 ---
    recent_start = date(2023, 4, 2)
    recent_end = date(2023, 6, 30)
    need_days = {d for d in all_days if recent_start <= d <= recent_end}
    # 也加入历年七月初与窗口日
    for y in (2020, 2021, 2022):
        for dd in range(1, 8):
            need_days.add(date(y, 7, dd))
    for dd in range(24, 31):
        need_days.add(date(2023, 6, dd))

    print("loading 批发价子集 …")
    wholesale = load_wholesale_for_days(need_days)
    print(f"wholesale pairs={len(wholesale)}")

    # 品类日：销量加权均售价 vs 销量加权均批发价 → 加成率
    markup_rows = []
    day_to_skus: dict[date, list[str]] = defaultdict(list)
    for (d, code), q in S["day_sku_qty"].items():
        if recent_start <= d <= recent_end and q != 0:
            day_to_skus[d].append(code)

    for d, codes in day_to_skus.items():
        agg = {c: {"q": 0.0, "rev": 0.0, "cost": 0.0, "q_cost": 0.0} for c in CAT_ORDER}
        for code in codes:
            q = S["day_sku_qty"][(d, code)]
            rev = S["day_sku_rev"][(d, code)]
            cat = sku2cat[code]
            wp = wholesale.get((d, code))
            agg[cat]["q"] += q
            agg[cat]["rev"] += rev
            if wp is not None and q > 0:
                agg[cat]["cost"] += wp * q
                agg[cat]["q_cost"] += q
        for cat in CAT_ORDER:
            a = agg[cat]
            if a["q"] <= 0 or a["q_cost"] <= 0:
                continue
            avg_sell = a["rev"] / a["q"]
            avg_cost = a["cost"] / a["q_cost"]
            if avg_cost <= 0:
                continue
            markup = (avg_sell - avg_cost) / avg_cost
            markup_rows.append(
                {
                    "日期": d,
                    "品类": cat,
                    "净销量_kg": a["q"],
                    "加权均售价": avg_sell,
                    "加权均批发价": avg_cost,
                    "加成率": markup,
                    "星期": d.weekday(),
                }
            )
    markup_df = pd.DataFrame(markup_rows)
    markup_df.to_csv(OUT_DATA / "近90品类日_量加成.csv", index=False, encoding="utf-8-sig")

    # 各品类：销量 vs 加成率 Pearson；分位
    sum_rows = []
    for cat in CAT_ORDER:
        sub = markup_df[markup_df["品类"] == cat]
        r = pearson(sub["加成率"].to_numpy(), sub["净销量_kg"].to_numpy())
        sum_rows.append(
            {
                "品类": cat,
                "n_日": len(sub),
                "加成率_P25": sub["加成率"].quantile(0.25),
                "加成率_P50": sub["加成率"].quantile(0.5),
                "加成率_P75": sub["加成率"].quantile(0.75),
                "pearson_加成率_vs_销量": r,
                "日均销量_kg": sub["净销量_kg"].mean(),
            }
        )
    sum_df = pd.DataFrame(sum_rows)
    sum_df.to_csv(OUT_DATA / "近90量加成摘要.csv", index=False, encoding="utf-8-sig")
    print(sum_df.to_string(index=False))

    fig, axes = plt.subplots(2, 3, figsize=(11, 7), sharex=False)
    for ax, cat in zip(axes.ravel(), CAT_ORDER):
        sub = markup_df[markup_df["品类"] == cat]
        ax.scatter(sub["加成率"], sub["净销量_kg"], s=12, alpha=0.55)
        ax.set_title(cat)
        ax.set_xlabel("加成率")
        ax.set_ylabel("日净销量kg")
    fig.suptitle("近90日：品类日加成率 vs 净销量")
    fig.tight_layout()
    fig.savefig(OUT_FIG / "量加成散点_近90.png", dpi=120)
    plt.close(fig)

    # --- 历年 7/1–7/7 品类日均 ---
    july_rows = []
    for y in (2020, 2021, 2022):
        for dd in range(1, 8):
            d = date(y, 7, dd)
            for cat in CAT_ORDER:
                july_rows.append(
                    {
                        "年": y,
                        "日期": d,
                        "星期": d.weekday(),
                        "品类": cat,
                        "净销量_kg": S["july_hist"].get((y, d, cat), 0.0),
                    }
                )
    july_df = pd.DataFrame(july_rows)
    july_df.to_csv(OUT_DATA / "历年7月初品类日销量.csv", index=False, encoding="utf-8-sig")
    july_sum = (
        july_df.groupby(["年", "品类"], as_index=False)["净销量_kg"]
        .agg(周合计="sum", 日均="mean")
        .sort_values(["品类", "年"])
    )
    july_sum.to_csv(OUT_DATA / "历年7月初周汇总.csv", index=False, encoding="utf-8-sig")
    print(july_sum.to_string(index=False))

    # 2023-07-01～07 星期（已知）与近90同星期日均对照
    target_week = [date(2023, 7, d) for d in range(1, 8)]
    wd_map = {i: [] for i in range(7)}
    for _, row in panel[
        (panel["日期"] >= recent_start) & (panel["日期"] <= recent_end)
    ].iterrows():
        wd_map[row["星期"]].append((row["品类"], row["净销量_kg"]))
    wd_mean = {}
    for wd in range(7):
        by_cat = defaultdict(list)
        for cat, q in wd_map[wd]:
            by_cat[cat].append(q)
        for cat in CAT_ORDER:
            wd_mean[(wd, cat)] = float(np.mean(by_cat[cat])) if by_cat[cat] else float("nan")

    forecast_rows = []
    for d in target_week:
        for cat in CAT_ORDER:
            forecast_rows.append(
                {
                    "目标日": d,
                    "星期": d.weekday(),
                    "品类": cat,
                    "近90同星期日均_kg": wd_mean[(d.weekday(), cat)],
                }
            )
    forecast_df = pd.DataFrame(forecast_rows)
    forecast_df.to_csv(OUT_DATA / "2023-07周_近90同星期基线.csv", index=False, encoding="utf-8-sig")
    print(forecast_df.groupby("品类")["近90同星期日均_kg"].mean())

    # --- Q3 窗口 ---
    win_codes = [c for c, q in S["win_sku_qty"].items() if q != 0]
    win_rows = []
    for code in win_codes:
        q7 = S["win_sku_qty"][code]
        days = len(S["win_sku_days"][code])
        daily = q7 / 7.0
        loss = loss_sku.loc[loss_sku["单品编码"] == code, "单品损耗率pct"]
        loss_v = float(loss.iloc[0]) if len(loss) else float("nan")
        # 窗口内批发价中位
        wps = [wholesale[(d, code)] for d in sorted(S["win_sku_days"][code]) if (d, code) in wholesale]
        win_rows.append(
            {
                "单品编码": code,
                "单品名称": sku2name.get(code, ""),
                "品类": sku2cat.get(code, ""),
                "窗口净销量_kg": q7,
                "销售天数": days,
                "窗口日均_kg": daily,
                "日均是否>=2.5": daily >= 2.5,
                "单品损耗率pct": loss_v,
                "窗口批发价中位": float(np.median(wps)) if wps else float("nan"),
            }
        )
    win_df = pd.DataFrame(win_rows).sort_values(["品类", "窗口净销量_kg"], ascending=[True, False])
    win_df.to_csv(OUT_DATA / "Q3窗口SKU画像.csv", index=False, encoding="utf-8-sig")
    n_win = len(win_df)
    n_ge25 = int(win_df["日均是否>=2.5"].sum())
    print(f"窗口SKU={n_win}, 日均>=2.5kg的={n_ge25}")
    print(win_df.groupby("品类").agg(SKU数=("单品编码", "count"), 日均过2点5=("日均是否>=2.5", "sum")))

    # 若只选日均>=2.5：品类覆盖与个数
    hard = win_df[win_df["日均是否>=2.5"]]
    print("硬过2.5按品类:\n", hard.groupby("品类").size())

    # 打折销量占比（近90）
    disc_rows = []
    for cat in CAT_ORDER:
        sub = panel[(panel["日期"] >= recent_start) & (panel["日期"] <= recent_end) & (panel["品类"] == cat)]
        tot = sub["净销量_kg"].sum()
        # 打折销量含负？净销量口径下打折行已在 disc 累计（含退货打折极少）
        disc = sub["打折销量_kg"].sum()
        disc_rows.append({"品类": cat, "近90净销量": tot, "近90打折销量": disc, "打折销量占比": disc / tot if tot else np.nan})
    disc_df = pd.DataFrame(disc_rows)
    disc_df.to_csv(OUT_DATA / "近90打折销量占比.csv", index=False, encoding="utf-8-sig")
    print(disc_df.to_string(index=False))

    # 损耗口径对照
    loss_cmp = loss_cat.copy()
    simple = (
        loss_sku.merge(catalog[["单品编码", "分类名称"]], on="单品编码")
        .groupby("分类名称")["单品损耗率pct"]
        .mean()
        .rename("单品简单均值")
    )
    loss_cmp = loss_cmp.merge(simple, left_on="分类名称", right_index=True)
    loss_cmp.to_csv(OUT_DATA / "损耗口径对照.csv", index=False, encoding="utf-8-sig")

    # 品类日销量近90箱线（相对窗口）
    fig, ax = plt.subplots(figsize=(8, 4.5))
    data = [
        panel[(panel["日期"] >= recent_start) & (panel["日期"] <= recent_end) & (panel["品类"] == c)][
            "净销量_kg"
        ].to_numpy()
        for c in CAT_ORDER
    ]
    ax.boxplot(data, tick_labels=CAT_ORDER, showfliers=False)
    ax.set_title("近90日品类日净销量（去极端）")
    ax.set_ylabel("kg")
    fig.tight_layout()
    fig.savefig(OUT_FIG / "近90品类日销量箱线.png", dpi=120)
    plt.close(fig)

    # 写摘要 json 友好文本
    summary = {
        "窗口SKU数": n_win,
        "日均>=2.5kg_SKU数": n_ge25,
        "2023-07-01_星期": "周六",
        "近365花叶-辣椒相关": float(c365.loc["花叶类", "辣椒类"]),
        "近365花叶-食用菌相关": float(c365.loc["花叶类", "食用菌"]),
        "近365花菜-水生相关": float(c365.loc["花菜类", "水生根茎类"]),
    }
    pd.Series(summary).to_json(OUT_DATA / "摘要.json", force_ascii=False, indent=2)
    print("DONE", summary)


if __name__ == "__main__":
    main()
