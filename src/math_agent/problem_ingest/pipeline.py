"""题面文件解析流水线：文本抽取 → 乱码检测 → 可选视觉回退 → MD 落盘。"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from math_agent.problem_ingest.garble import assess_garble
from math_agent.problem_ingest.markdown_out import write_parsed_markdown
from math_agent.problem_ingest.pdf_text import extract_pdf_text, sanitize_text


@dataclass
class ParseQuality:
    ok: bool
    garble_ratio: float
    method: str
    warnings: list[str] = field(default_factory=list)
    pages: int = 0
    needs_vision: bool = False

    def to_dict(self) -> dict:
        return {
            "ok": self.ok,
            "garbleRatio": self.garble_ratio,
            "method": self.method,
            "warnings": list(self.warnings),
            "pages": self.pages,
            "needsVision": self.needs_vision,
        }


@dataclass
class ProblemParseResult:
    text: str
    quality: ParseQuality
    parsed_md_path: Path | None = None
    total_pages: int = 0

    def to_summary_dict(self, *, text_limit: int = 5000) -> dict:
        """兼容 extract_file_meta 的 summary 形状，并附带质量字段。"""
        return {
            "text_excerpt": self.text[:text_limit],
            "total_pages": self.total_pages,
            "parse_quality": self.quality.to_dict(),
            "parsed_md": str(self.parsed_md_path).replace("\\", "/") if self.parsed_md_path else "",
        }


def _read_docx(path: Path) -> tuple[str, dict]:
    from docx import Document

    doc = Document(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    text = "\n".join(paragraphs)
    meta = {"paragraphs": len(paragraphs), "tables": len(doc.tables)}
    return text, meta


def _read_plain(path: Path) -> str:
    return sanitize_text(path.read_text(encoding="utf-8", errors="ignore"))


def write_problem_source_md(out_dir: Path | str, body: str, *, source: str = "user_confirmed") -> Path:
    """将启动时最终题面写入 run 输出目录的 problem_source.md。"""
    return write_parsed_markdown(
        Path(out_dir) / "problem_source.md",
        body,
        source=source,
        method="user_confirmed",
        garble_score=0.0,
        pages=0,
        warnings=[],
    )


def parse_problem_file(
    path: Path | str,
    *,
    enable_vision_fallback: bool = True,
    write_md: bool = True,
    md_filename: str = "problem_parsed.md",
    complete_fn=None,
) -> ProblemParseResult:
    """解析题面文件；PDF 乱码超阈值时可视觉回退。"""
    path = Path(path)
    suffix = path.suffix.lower()
    warnings: list[str] = []
    method = "direct"
    pages = 0
    text = ""

    if suffix == ".pdf":
        text, pages = extract_pdf_text(path, include_tables=True)
        method = "text"
        report = assess_garble(text)
        warnings = list(report.warnings)
        garble_ratio = report.garble_ratio

        needs_vision = report.should_vision_fallback
        if enable_vision_fallback and needs_vision:
            try:
                from math_agent.problem_ingest.vision import transcribe_pdf_with_vision

                vision_text = transcribe_pdf_with_vision(path, complete_fn=complete_fn)
                if vision_text.strip():
                    text = vision_text
                    method = "vision"
                    needs_vision = False
                    vision_report = assess_garble(text)
                    warnings = list(
                        dict.fromkeys(warnings + ["vision_fallback_used"] + vision_report.warnings)
                    )
                    garble_ratio = vision_report.garble_ratio
                else:
                    method = "text_fallback"
                    warnings = list(dict.fromkeys(warnings + ["vision_empty_fallback"]))
            except Exception as exc:  # noqa: BLE001 — 视觉失败不阻断上传
                method = "text_fallback"
                warnings = list(dict.fromkeys(warnings + [f"vision_failed:{type(exc).__name__}"]))

        if method == "vision":
            quality_ok = True
        elif method == "text_fallback":
            quality_ok = False
        else:
            # 文本层可用但建议视觉时：ok=False 便于 UI 提示，needs_vision=True
            quality_ok = report.ok and not needs_vision
        quality = ParseQuality(
            ok=quality_ok,
            garble_ratio=garble_ratio,
            method=method,
            warnings=warnings,
            pages=pages,
            needs_vision=needs_vision and method == "text",
        )
    elif suffix == ".docx":
        text, _meta = _read_docx(path)
        method = "docx"
        report = assess_garble(text)
        quality = ParseQuality(
            ok=report.ok,
            garble_ratio=report.garble_ratio,
            method=method,
            warnings=list(report.warnings),
            pages=0,
        )
    elif suffix in {".md", ".txt"}:
        text = _read_plain(path)
        method = "direct"
        quality = ParseQuality(ok=True, garble_ratio=0.0, method=method, warnings=[], pages=0)
    else:
        raise ValueError(f"unsupported problem file type: {suffix}")

    md_path = None
    if write_md:
        md_path = write_parsed_markdown(
            path.parent / md_filename,
            text,
            source=path.name,
            method=quality.method,
            garble_score=quality.garble_ratio,
            pages=quality.pages,
            warnings=quality.warnings,
        )

    return ProblemParseResult(
        text=text,
        quality=quality,
        parsed_md_path=md_path,
        total_pages=pages,
    )


def apply_vision_transcription(
    path: Path | str,
    *,
    write_md: bool = True,
    md_filename: str = "problem_parsed.md",
    complete_fn=None,
    progress_cb=None,
) -> ProblemParseResult:
    """对已上传的题面 PDF 执行视觉转写，并刷新 problem_parsed.md。"""
    path = Path(path)
    if path.suffix.lower() != ".pdf":
        raise ValueError("vision transcription only supports PDF problem files")

    from math_agent.problem_ingest.vision import transcribe_pdf_with_vision

    # 先保留文本层，便于失败回退
    text_layer, pages = extract_pdf_text(path, include_tables=True)
    text_report = assess_garble(text_layer)
    warnings = list(text_report.warnings)
    try:
        vision_text = transcribe_pdf_with_vision(
            path, complete_fn=complete_fn, progress_cb=progress_cb,
        )
    except Exception as exc:  # noqa: BLE001
        quality = ParseQuality(
            ok=False,
            garble_ratio=text_report.garble_ratio,
            method="text_fallback",
            warnings=list(dict.fromkeys(warnings + [f"vision_failed:{type(exc).__name__}"])),
            pages=pages,
            needs_vision=False,
        )
        md_path = None
        if write_md:
            md_path = write_parsed_markdown(
                path.parent / md_filename,
                text_layer,
                source=path.name,
                method=quality.method,
                garble_score=quality.garble_ratio,
                pages=pages,
                warnings=quality.warnings,
            )
        return ProblemParseResult(
            text=text_layer, quality=quality, parsed_md_path=md_path, total_pages=pages,
        )

    if not vision_text.strip():
        quality = ParseQuality(
            ok=False,
            garble_ratio=text_report.garble_ratio,
            method="text_fallback",
            warnings=list(dict.fromkeys(warnings + ["vision_empty_fallback"])),
            pages=pages,
            needs_vision=False,
        )
        text = text_layer
    else:
        vision_report = assess_garble(vision_text)
        quality = ParseQuality(
            ok=True,
            garble_ratio=vision_report.garble_ratio,
            method="vision",
            warnings=list(dict.fromkeys(warnings + ["vision_fallback_used"] + vision_report.warnings)),
            pages=pages,
            needs_vision=False,
        )
        text = vision_text

    md_path = None
    if write_md:
        md_path = write_parsed_markdown(
            path.parent / md_filename,
            text,
            source=path.name,
            method=quality.method,
            garble_score=quality.garble_ratio,
            pages=pages,
            warnings=quality.warnings,
        )
    return ProblemParseResult(
        text=text, quality=quality, parsed_md_path=md_path, total_pages=pages,
    )
